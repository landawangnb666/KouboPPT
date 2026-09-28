"""从课时内容规格构建精美 pptx（本地渲染，只在"自动选风格"时问一次 AI）。

版式：封面 / 目录 / 章节过渡页 / 内容页。
风格 = 配色 + 版式骨架（chrome 页眉骨架 × cover 封面骨架），见 THEMES；
骨架只用色块与文字摆位实现，不引外部模板，任何风格都能安全渲染。
含公式的行整行渲染为高清图片（matplotlib），从根上避免乱码；万一本地渲不出来，
可注入 formula_fallback 兜底（见 formula_fallback 模块：AI 改写重渲 → 原书扫描页裁图 → 才降级文字）。
口播稿写入幻灯片备注，直接对接现有"备注"文稿来源的视频流水线。

"自动"模式（theme="auto"）按教材书名让 AI 从风格库里挑一套，见 resolve_style。

引擎接口 SlideEngine.build 预留：后续可替换/新增 PPTAgent 等实现。
"""
from __future__ import annotations

import json
from dataclasses import dataclass, field
from pathlib import Path
from typing import Protocol

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from . import formula, xmlsafe
from .llm import LLMClient

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
BODY_TOP = Inches(1.42)
BODY_LEFT = Inches(0.85)
BODY_RIGHT = Inches(12.48)          # 内容区右边界
MAX_BODY_BOTTOM = Inches(7.05)      # 超过则换"续页"


@dataclass
class SlideSpec:
    title: str
    lines: list[str] = field(default_factory=list)   # 要点，可含 $..$/$$..$$
    script: str = ""                                 # 口播稿（进备注）


@dataclass
class SectionSpec:
    title: str
    slides: list[SlideSpec] = field(default_factory=list)


@dataclass
class LessonSpec:
    title: str                          # 如 "第3章 微分中值定理"
    subtitle: str                       # 教材名
    sections: list[SectionSpec] = field(default_factory=list)
    theme: str = "stem_blue"


# 风格 = 配色（primary/accent/dark/text/muted/band）+ 骨架（chrome/cover）。
# chrome：页眉骨架 band（默认）/ sidebar / block / minimal
# cover ：封面骨架 full（默认）/ split / block / minimal
# label/note 只给界面下拉与 AI 选风格用，不参与渲染；key 一旦发布就不要改名（写进配置与缓存）。
THEMES: dict[str, dict[str, str]] = {
    "stem_blue": {"primary": "1F4E79", "accent": "2E9BD6", "dark": "16324A",
                  "text": "2B2B2B", "muted": "8A99A8", "band": "F2F6FA",
                  "chrome": "band", "cover": "full",
                  "label": "学术蓝（数学/物理）",
                  "note": "深蓝底、天蓝点缀，稳重学术，适合数学、物理等理科"},
    "engineering_gray": {"primary": "37474F", "accent": "F26419", "dark": "263238",
                         "text": "2B2B2B", "muted": "90A4AE", "band": "F5F5F5",
                         "chrome": "band", "cover": "full",
                         "label": "工业灰橙（机械/工科）",
                         "note": "灰蓝配安全橙，硬朗实用，适合机械、工程、自动化"},
    "science_green": {"primary": "1E5945", "accent": "8BC34A", "dark": "143B2E",
                      "text": "2B2B2B", "muted": "8FAA9E", "band": "F1F7F3",
                      "chrome": "band", "cover": "full",
                      "label": "理科墨绿（化学/生物）",
                      "note": "墨绿配草绿，自然清爽，适合化学、生物、环境"},
    "humanities_warm": {"primary": "6B4226", "accent": "C4762F", "dark": "3A2A1E",
                        "text": "2B2B2B", "muted": "97836F", "band": "FAF6F0",
                        "chrome": "minimal", "cover": "block",
                        "label": "人文暖棕（文科/历史）",
                        "note": "米白底、赭棕点缀，杂志感留白，适合语文、历史、政治等人文社科"},
    "medical_teal": {"primary": "0F5C63", "accent": "2BA89B", "dark": "103F44",
                     "text": "2B2B2B", "muted": "86A6A4", "band": "EFF6F5",
                     "chrome": "sidebar", "cover": "split",
                     "label": "医学青蓝（医学/护理）",
                     "note": "青色侧栏加留白，洁净克制，适合医学、护理、药学、生物医学"},
    "tech_indigo": {"primary": "2B3A8F", "accent": "5C7CFA", "dark": "1B2350",
                    "text": "2B2B2B", "muted": "8B93B5", "band": "F1F3FB",
                    "chrome": "block", "cover": "split",
                    "label": "科技靛蓝（计算机/电子）",
                    "note": "靛蓝圆角色块加亮色点缀，现代感强，适合计算机、电子、信息"},
    "mono_minimal": {"primary": "22262B", "accent": "C2410C", "dark": "14181C",
                     "text": "2B2B2B", "muted": "8A9096", "band": "F5F6F7",
                     "chrome": "minimal", "cover": "minimal",
                     "label": "素白极简（通用）",
                     "note": "纯白底、近黑标题、单一暖色点缀，克制通用，适合讲义、通识、复习提纲"},
}

DEFAULT_STYLE = "stem_blue"     # 兜底风格（识别不出学科时用）
AUTO_STYLE = "auto"             # 交给 AI 按书名选
STYLE_FILE = "PPT风格.json"      # 每本书的风格决定落盘在这里，重跑不再问 AI


def style_labels() -> dict[str, str]:
    """下拉框用：风格 key → 中文名（保持 THEMES 的注册顺序）。"""
    return {k: v.get("label", k) for k, v in THEMES.items()}


class SlideEngine(Protocol):
    def build(self, lesson: LessonSpec, out_path: Path) -> Path: ...


def _rgb(theme: dict, key: str) -> RGBColor:
    return RGBColor.from_string(theme[key])


def _set_run(run, size: float, color: RGBColor, bold: bool = False,
             name: str = "Microsoft YaHei") -> None:
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {})
        rPr.append(ea)
    ea.set("typeface", name)


def _put(p, text: str, size: float, color: RGBColor, bold: bool = False):
    """往段落写一段文字：先清洗 XML 非法字符，再统一设字体。"""
    r = p.add_run()
    r.text = xmlsafe.clean(text)
    _set_run(r, size, color, bold=bold)
    return r


def _textbox(slide, x, y, w, h):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    return tb, tf


def _rect(slide, x, y, w, h, color: RGBColor, shape=MSO_SHAPE.RECTANGLE):
    sp = slide.shapes.add_shape(shape, x, y, w, h)
    sp.fill.solid()
    sp.fill.fore_color.rgb = color
    sp.line.fill.background()
    sp.shadow.inherit = False
    return sp


def _px_size(png: Path) -> tuple[int, int]:
    from PIL import Image
    with Image.open(png) as im:
        return im.size


def _fit_emu(png: Path, max_w: Emu, scale: float = 1.0) -> tuple[Emu, Emu]:
    """PNG 按 300dpi 的原始物理尺寸缩放，宽度不超 max_w。"""
    w_px, h_px = _px_size(png)
    w = Emu(int(Emu(Inches(w_px / 300.0)) * scale))
    h = Emu(int(Emu(Inches(h_px / 300.0)) * scale))
    if w > max_w:
        h = Emu(int(h * max_w / w))
        w = max_w
    return w, h


class BuiltinEngine:
    """默认引擎：python-pptx 模板排版 + 公式图片化。

    formula_fallback：本地渲不出的公式的兜底回调，签名 ``fn(text, color, display) -> Path | None``
    ——返回一张图片就用它排进幻灯片（见 formula_fallback 模块：AI 改写重渲 → 扫图裁图），
    返回 None 才降级成可读文字。
    """

    def __init__(self, formula_cache: Path | None = None, formula_fallback=None):
        self.formula_cache = Path(formula_cache) if formula_cache else Path(".")
        self.formula_fallback = formula_fallback
        self.failed_formulas: list[str] = []      # 连兜底都救不回来的
        self.formula_rescued: list[str] = []      # 靠兜底图救回来的
        self.warnings: list[str] = []

    # ------------------------------------------------------------------
    def build(self, lesson: LessonSpec, out_path: Path) -> Path:
        theme = THEMES.get(lesson.theme, THEMES["stem_blue"])
        prs = Presentation()
        prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
        blank = prs.slide_layouts[6]

        self._cover(prs, blank, theme, lesson)
        if len(lesson.sections) > 1:
            self._toc(prs, blank, theme, lesson)
        for idx, sec in enumerate(lesson.sections, 1):
            if len(lesson.sections) > 1:
                self._section_page(prs, blank, theme, idx, sec)
            for sli in sec.slides:
                self._content(prs, blank, theme, sli, idx, sec.title)

        out_path = Path(out_path)
        out_path.parent.mkdir(parents=True, exist_ok=True)
        prs.save(out_path)
        return out_path

    # ------------------------------------------------------------------
    def _cover(self, prs, blank, theme, lesson: LessonSpec):
        """封面骨架：full（整幅深色，默认）/ split（左右分割）/ block（白底深色块）/ minimal（极简白）。"""
        s = prs.slides.add_slide(blank)
        variant = str(theme.get("cover", "full"))
        white = RGBColor.from_string("FFFFFF")
        accent = _rgb(theme, "accent")
        if variant == "split":
            w = Inches(7.5)
            _rect(s, 0, 0, w, SLIDE_H, _rgb(theme, "primary"))
            _rect(s, w, 0, Inches(0.07), SLIDE_H, accent)
            tb, tf = _textbox(s, Inches(0.85), Inches(2.35), Inches(6.1), Inches(2.6))
            _put(tf.paragraphs[0], lesson.title, 34, white, bold=True)
            tb2, tf2 = _textbox(s, Inches(0.85), Inches(4.95), Inches(6.1), Inches(0.9))
            _put(tf2.paragraphs[0], lesson.subtitle, 16, _rgb(theme, "muted"))
            _rect(s, w + Inches(0.95), Inches(2.7), Inches(3.5), Inches(0.07), accent)
        elif variant == "block":
            _rect(s, 0, Inches(4.15), SLIDE_W, SLIDE_H - Inches(4.15),
                  _rgb(theme, "primary"))
            _rect(s, 0, Inches(4.15), SLIDE_W, Inches(0.07), accent)
            tb, tf = _textbox(s, Inches(0.9), Inches(1.5), Inches(11.5), Inches(2.0))
            _put(tf.paragraphs[0], lesson.title, 40, _rgb(theme, "dark"), bold=True)
            tb2, tf2 = _textbox(s, Inches(0.9), Inches(3.32), Inches(11.5), Inches(0.7))
            _put(tf2.paragraphs[0], lesson.subtitle, 18, _rgb(theme, "muted"))
        elif variant == "minimal":
            _rect(s, Inches(0.9), Inches(2.2), Inches(0.95), Inches(0.09), accent)
            tb, tf = _textbox(s, Inches(0.9), Inches(2.6), Inches(11.5), Inches(2.2))
            _put(tf.paragraphs[0], lesson.title, 40, _rgb(theme, "dark"), bold=True)
            tb2, tf2 = _textbox(s, Inches(0.9), Inches(4.6), Inches(11.5), Inches(0.8))
            _put(tf2.paragraphs[0], lesson.subtitle, 18, _rgb(theme, "muted"))
        else:                                       # full（默认，保持原样）
            _rect(s, 0, 0, SLIDE_W, SLIDE_H, _rgb(theme, "primary"))
            _rect(s, 0, Inches(5.9), SLIDE_W, Inches(0.06), accent)
            tb, tf = _textbox(s, Inches(0.9), Inches(2.4), Inches(11.5), Inches(2.2))
            _put(tf.paragraphs[0], lesson.title, 40, white, bold=True)
            tb2, tf2 = _textbox(s, Inches(0.9), Inches(4.6), Inches(11.5), Inches(0.8))
            _put(tf2.paragraphs[0], lesson.subtitle, 18, _rgb(theme, "muted"))
        self._set_notes(s, f"这是《{lesson.title}》这节课。我们先来看本节的学习内容。")

    def _toc(self, prs, blank, theme, lesson: LessonSpec):
        s = prs.slides.add_slide(blank)
        self._page_chrome(s, theme, "课程目录", lesson.title)
        body = [f"{'壹贰叁肆伍陆柒捌'[i] if i < 8 else i + 1}、{sec.title}"
                for i, sec in enumerate(lesson.sections)]
        self._lines_area(s, theme, [f"# {b}" for b in body], Inches(1.5))

    def _section_page(self, prs, blank, theme, idx: int, sec: SectionSpec):
        """章节过渡页：跟随封面骨架（split 分栏 / minimal 白底，其余用整幅深色）。"""
        s = prs.slides.add_slide(blank)
        variant = str(theme.get("cover", "full"))
        accent = _rgb(theme, "accent")
        if variant == "split":
            w = Inches(7.5)
            _rect(s, 0, 0, w, SLIDE_H, _rgb(theme, "primary"))
            _rect(s, w, 0, Inches(0.07), SLIDE_H, accent)
            tb, tf = _textbox(s, Inches(0.85), Inches(2.3), Inches(6.1), Inches(1.1))
            _put(tf.paragraphs[0], f"{idx:02d}", 42, accent, bold=True)
            tb2, tf2 = _textbox(s, Inches(0.85), Inches(3.45), Inches(6.1), Inches(1.6))
            _put(tf2.paragraphs[0], sec.title, 30, RGBColor.from_string("FFFFFF"), bold=True)
        elif variant == "minimal":
            _rect(s, BODY_LEFT, Inches(2.5), Inches(0.95), Inches(0.09), accent)
            tb, tf = _textbox(s, Inches(0.9), Inches(2.85), Inches(11.5), Inches(1.1))
            _put(tf.paragraphs[0], f"{idx:02d}", 40, accent, bold=True)
            tb2, tf2 = _textbox(s, Inches(0.9), Inches(4.0), Inches(11.5), Inches(1.6))
            _put(tf2.paragraphs[0], sec.title, 30, _rgb(theme, "dark"), bold=True)
        else:                                       # full / block：整幅深色（原样）
            _rect(s, 0, 0, SLIDE_W, SLIDE_H, _rgb(theme, "primary"))
            tb, tf = _textbox(s, Inches(0.9), Inches(2.2), Inches(11.5), Inches(1.2))
            _put(tf.paragraphs[0], f"{idx:02d}", 44, accent, bold=True)
            tb2, tf2 = _textbox(s, Inches(0.9), Inches(3.4), Inches(11.5), Inches(1.6))
            _put(tf2.paragraphs[0], sec.title, 32, RGBColor.from_string("FFFFFF"), bold=True)
        self._set_notes(s, f"接下来进入下一个部分：{sec.title}。")

    def _page_chrome(self, s, theme, title: str, kicker: str):
        """页眉骨架：band（顶部色条，默认）/ sidebar（左侧色栏）/ block（色块标题）/ minimal（极简）。

        四种骨架的标题区都控制在 1.4 英寸以内，正文一律从 BODY_TOP 起，互不干扰。
        """
        variant = str(theme.get("chrome", "band"))
        accent, dark, muted = (_rgb(theme, "accent"), _rgb(theme, "dark"),
                               _rgb(theme, "muted"))
        if variant == "sidebar":
            band = Inches(0.5)
            _rect(s, 0, 0, band, SLIDE_H, _rgb(theme, "primary"))
            _rect(s, band, 0, Inches(0.055), SLIDE_H, accent)
            left = band + Inches(0.45)
            width = SLIDE_W - left - Inches(0.5)
            tb, tf = _textbox(s, left, Inches(0.44), width, Inches(0.9))
            _put(tf.paragraphs[0], title, 26, dark, bold=True)
            tb2, tf2 = _textbox(s, left, Inches(1.1), width, Inches(0.4))
            _put(tf2.paragraphs[0], kicker, 12, muted)
            _rect(s, left, Inches(1.38), Inches(1.1), Inches(0.045), accent)
        elif variant == "block":
            sp = _rect(s, BODY_LEFT, Inches(0.3), Inches(7.6), Inches(0.84),
                       _rgb(theme, "primary"), MSO_SHAPE.ROUNDED_RECTANGLE)
            try:
                sp.adjustments[0] = 0.16
            except Exception:                      # 个别模板无调整柄，退化成直角也无妨
                pass
            tf = sp.text_frame
            tf.word_wrap = True
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf.margin_left, tf.margin_right = Inches(0.24), Inches(0.2)
            size = 24 if len(title) <= 16 else (20 if len(title) <= 24 else 17)
            _put(tf.paragraphs[0], title, size, RGBColor.from_string("FFFFFF"), bold=True)
            tf.paragraphs[0].alignment = PP_ALIGN.LEFT   # 自选图形的文字默认居中，这里跟正文左对齐
            tb2, tf2 = _textbox(s, BODY_LEFT, Inches(1.24), Inches(11.6), Inches(0.4))
            _put(tf2.paragraphs[0], kicker, 11, muted)
        elif variant == "minimal":
            tb, tf = _textbox(s, BODY_LEFT, Inches(0.46), Inches(11.6), Inches(0.95))
            _put(tf.paragraphs[0], title, 29, dark, bold=True)
            tb2, tf2 = _textbox(s, BODY_LEFT, Inches(1.14), Inches(11.6), Inches(0.36))
            _put(tf2.paragraphs[0], kicker, 11, muted)
            _rect(s, BODY_LEFT, Inches(1.4), BODY_RIGHT - BODY_LEFT, Inches(0.022), muted)
        else:                                       # band（默认，保持原样）
            _rect(s, 0, 0, SLIDE_W, Inches(0.12), accent)
            tb, tf = _textbox(s, BODY_LEFT, Inches(0.42), Inches(11.6), Inches(0.9))
            _put(tf.paragraphs[0], title, 26, dark, bold=True)
            tb2, tf2 = _textbox(s, BODY_LEFT, Inches(1.08), Inches(11.6), Inches(0.4))
            _put(tf2.paragraphs[0], kicker, 12, muted)
            _rect(s, BODY_LEFT, Inches(1.38), Inches(1.1), Inches(0.045), accent)

    # ------------------------------------------------------------------
    def _content(self, prs, blank, theme, sli: SlideSpec, sec_idx: int,
                 sec_title: str, overflow: list | None = None, cont: int = 0):
        lines = overflow if overflow is not None else list(sli.lines)
        title = sli.title if overflow is None else f"{sli.title}（续）"
        s = prs.slides.add_slide(blank)
        self._page_chrome(s, theme, title, f"{sec_idx:02d} · {sec_title}")
        next_lines = self._lines_area(s, theme, lines, BODY_TOP)
        if overflow is None:
            self._set_notes(s, sli.script)
        # 续页故意不写备注：视频流水线把备注当口播稿，写什么都会被朗读出来
        if not next_lines:
            return
        if cont >= 1:
            # 整页也放不下（典型是又窄又高的公式图）：只允许一层续页，多余内容
            # 强制放在新一页上并记警告，避免无限递归
            self.warnings.append(
                f"《{sli.title}》的内容超过两页也排不完，最后一部分可能超出页面，请手工检查")
            extra = prs.slides.add_slide(blank)
            self._page_chrome(extra, theme, f"{sli.title}（续）",
                              f"{sec_idx:02d} · {sec_title}")
            self._lines_area(extra, theme, next_lines, BODY_TOP, force=True)
            return
        self._content(prs, blank, theme, sli, sec_idx, sec_title, next_lines, cont + 1)

    def _lines_area(self, s, theme, lines: list[str], top: Emu,
                    force: bool = False) -> list[str]:
        """自上而下排版；排不下的行返回给调用方做续页。行前缀 '# ' 表示无符号行。
        force=True 时不再判断溢出，全部硬排（最后手段，配合警告使用）。"""
        y = top + Inches(0.18)
        i = 0
        while i < len(lines):
            raw = lines[i].strip()
            if not raw:
                i += 1
                continue
            plain = raw.startswith("# ")
            text = raw[2:] if plain else raw
            h_used, item = self._render_line(s, theme, text, y, plain, force)
            if item is None and h_used is None:      # 溢出信号
                return lines[i:]
            est_h = Inches(0.52) if h_used is None else h_used
            y += est_h + Inches(0.14)
            i += 1
        return []

    def _render_line(self, s, theme, text: str, y: Emu, plain: bool,
                     force: bool = False):
        """渲染一行。返回 (占用高度或None, 需要挪到 y 的形状或None)。
        None,None 表示该行放不下（溢出）。"""
        remaining = MAX_BODY_BOTTOM - y
        text = formula.normalize_delims(text)
        display_math = text.startswith("$$") and text.endswith("$$")
        if formula.has_math(text) or display_math:
            got = self._formula_image(text, theme, display_math)
            if got is not None:
                indent = Emu(0) if plain else Inches(0.4)
                max_w = BODY_RIGHT - BODY_LEFT - indent
                w, h = _fit_emu(got, max_w)
                while h > remaining and w > Inches(2.5):   # 放不下先缩一点
                    shrink = int(w * 0.85) or 1
                    h = int(h * shrink / w)
                    w = shrink
                if h > remaining and not force:
                    return None, None
                x = BODY_LEFT + (Emu(int((BODY_RIGHT - BODY_LEFT - w) / 2))
                                 if display_math else indent)
                pic = s.shapes.add_picture(str(got), x, y, width=w)
                return h, pic
            self.failed_formulas.append(text)
        # 普通文字行（含公式渲染失败的降级：转成可读文本，别甩 LaTeX 源码）
        fallback = formula.to_readable(text)
        bullet = "" if plain else "• "
        tb, tf = _textbox(s, BODY_LEFT, y, BODY_RIGHT - BODY_LEFT - Inches(0.4), Inches(1.2))
        p = tf.paragraphs[0]
        _put(p, bullet + fallback, 18 if not plain else 20, _rgb(theme, "text"), bold=plain)
        est = Inches(0.5) if len(fallback) < 40 else Inches(0.95)
        if y + est > MAX_BODY_BOTTOM and not force:
            tb._element.getparent().remove(tb._element)
            return None, None
        return est, tb

    def _formula_image(self, text: str, theme, display_math: bool) -> Path | None:
        """这一行的公式图：先本地渲；渲不出来才问兜底回调（AI 改写重渲 / 扫图裁图）。

        兜底只在失败时发生，正常路径一点不受影响；回调自己抛异常也不能拦下整节课。
        """
        color = "#" + theme["text"]
        png = formula.render_dir_for(self.formula_cache, text, 20, color, display_math)
        got = formula.render(text, png, fontsize=20, color=color, display=display_math)
        if got is not None or self.formula_fallback is None:
            return got
        try:
            rescued = self.formula_fallback(text, color, display_math)
        except Exception as exc:                    # noqa: BLE001 兜底不能成为新的失败点
            self.warnings.append(
                f"公式兜底失败（{type(exc).__name__}: {exc}）：{text[:60]}")
            return None
        if rescued is not None:
            self.formula_rescued.append(text)
        return rescued

    # ------------------------------------------------------------------
    def _set_notes(self, s, text: str):
        text = xmlsafe.clean(text)          # 口播稿也要清洗：备注同样是 OpenXML
        if not text:
            return
        s.notes_slide.notes_text_frame.text = text


def build(lesson: LessonSpec, out_path: Path, theme: str = DEFAULT_STYLE,
          formula_cache: Path | None = None, formula_fallback=None) -> Path:
    """渲染一节课的 pptx。theme 传风格 key；"auto" 在此兜底为默认风格
    （正常路径由 courseware/GUI 先调 resolve_style 解析成具体风格）。"""
    if theme in THEMES:
        lesson.theme = theme
    elif lesson.theme not in THEMES:
        lesson.theme = DEFAULT_STYLE
    engine = BuiltinEngine(formula_cache=formula_cache or Path(out_path).parent / ".eq_cache",
                           formula_fallback=formula_fallback)
    return engine.build(lesson, out_path)


# ---------------------------------------------------------------- 风格：AI 按书名选
_STYLE_SYS = ("你是课件视觉设计师。根据教材名称判断学科与气质，"
              "从给定风格清单里挑出最合适的几套，并按推荐度排序。")


def _style_catalog() -> str:
    return "\n".join(f"- {k}：{v.get('label', k)}；{v.get('note', '')}"
                     for k, v in THEMES.items())


def pick_style_candidates(llm: LLMClient, book_title: str, n: int = 3,
                          raw_dir: Path | None = None) -> list[dict]:
    """让 AI 按书名挑风格，返回 [{"key","reason"}, ...]（已过滤非法 key、已去重）。

    只让 AI 在 THEMES 里**选**，不允许它自造配色——这是质量下限的保证。
    解析不出任何合法项时返回空列表，由调用方回落默认风格。
    """
    data = llm.chat_json(
        _STYLE_SYS,
        f"教材名称：《{book_title}》\n\n可选风格：\n{_style_catalog()}\n\n"
        f"请挑出最合适的 {max(1, min(n, len(THEMES)))} 套（按推荐度从高到低），"
        "每套给一句中文理由（说明为什么适合这本书）。\n"
        '输出 JSON：{"styles":[{"key":"风格key","reason":"理由"}]}',
        max_tokens=800, retries=1, raw_dir=raw_dir, tag="PPT风格选择")
    items = data.get("styles") if isinstance(data, dict) else data
    out: list[dict] = []
    for it in items or []:
        key = str((it or {}).get("key", "")).strip()
        if key in THEMES and all(o["key"] != key for o in out):
            out.append({"key": key, "reason": str(it.get("reason", "")).strip()})
    return out[:max(1, n)]


def load_style_choice(book_dir: Path) -> dict | None:
    """读这本书已定的风格（缓存）。文件坏了/风格已下线都当没有。"""
    try:
        data = json.loads((Path(book_dir) / STYLE_FILE).read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return None
    key = str((data or {}).get("key", "")).strip()
    if key not in THEMES:
        return None
    return {"key": key, "reason": str(data.get("reason", ""))}


def save_style_choice(book_dir: Path, key: str, reason: str = "") -> None:
    """把风格决定落盘（书目录下），重跑不再问 AI、全书各节风格一致。"""
    if key not in THEMES:
        return
    try:
        d = Path(book_dir)
        d.mkdir(parents=True, exist_ok=True)
        (d / STYLE_FILE).write_text(json.dumps(
            {"key": key, "label": THEMES[key].get("label", key), "reason": reason},
            ensure_ascii=False, indent=2), encoding="utf-8")
    except OSError:
        pass


def resolve_style(llm: LLMClient, book_title: str, book_dir: Path, log=print,
                  raw_dir: Path | None = None) -> str:
    """自动模式：命中缓存直接用；否则问一次 AI 取推荐度最高的风格；失败回落默认。"""
    cached = load_style_choice(book_dir)
    if cached:
        log(f"  PPT 风格：沿用已选「{THEMES[cached['key']].get('label', cached['key'])}」"
            + (f"（{cached['reason']}）" if cached["reason"] else ""))
        return cached["key"]
    try:
        cands = pick_style_candidates(llm, book_title, raw_dir=raw_dir)
    except Exception as exc:                       # noqa: BLE001  选风格失败不该拖垮主流程
        log(f"  ⚠ AI 选风格失败（{exc}），改用默认风格")
        cands = []
    if not cands:
        log(f"  PPT 风格：默认「{THEMES[DEFAULT_STYLE].get('label', DEFAULT_STYLE)}」")
        return DEFAULT_STYLE
    pick = cands[0]
    label = THEMES[pick["key"]].get("label", pick["key"])
    log(f"  PPT 风格：AI 按书名选定「{label}」"
        + (f"——{pick['reason']}" if pick["reason"] else ""))
    save_style_choice(book_dir, pick["key"], (pick["reason"] + "（AI 自动选定）").strip())
    return pick["key"]


# ---------------------------------------------------------------- 风格预览缩略图
def style_preview_image(key: str, width: int = 420, height: int = 250):
    """离线画一张风格预览图（左封面骨架 + 右内容页骨架），不调任何接口。

    只按骨架规则摆色块，够在候选里分辨布局与配色即可（不追求与成片逐像素一致）。
    """
    from PIL import Image, ImageDraw

    th = THEMES.get(key) or THEMES[DEFAULT_STYLE]

    def col(name: str, dflt: str = "FFFFFF") -> str:
        return "#" + str(th.get(name, dflt)).lstrip("#")

    primary, accent = col("primary", "1F4E79"), col("accent", "2E9BD6")
    dark, muted, page = col("dark", "16324A"), col("muted", "8A99A8"), "#FFFFFF"
    white = "#FFFFFF"
    line = "#D7DCE3"
    img = Image.new("RGB", (width, height), "#E9EDF3")
    d = ImageDraw.Draw(img)
    pad, gap = 14, 18
    cw = (width - pad * 2 - gap) // 2
    ch = height - pad * 2

    def bar(x0, y0, x1, y1, color):
        d.rectangle((x0, y0, x1, y1), fill=color)

    # ---- 左：封面骨架
    x0, y0, x1, y1 = pad, pad, pad + cw, pad + ch
    d.rounded_rectangle((x0, y0, x1, y1), radius=6, fill=page)
    cv = str(th.get("cover", "full"))
    if cv == "split":
        split = x0 + int(cw * 0.58)
        bar(x0, y0, split, y1, primary)
        bar(split, y0, split + 3, y1, accent)
        bar(x0 + 16, y0 + int(ch * 0.34), x0 + int(cw * 0.42), y0 + int(ch * 0.34) + 8, white)
        bar(x0 + 16, y0 + int(ch * 0.34) + 16, x0 + int(cw * 0.34), y0 + int(ch * 0.34) + 22, white)
        bar(x0 + 16, y0 + int(ch * 0.34) + 34, x0 + int(cw * 0.28), y0 + int(ch * 0.34) + 39, muted)
        bar(split + 24, y0 + int(ch * 0.52), x1 - 24, y0 + int(ch * 0.52) + 4, accent)
    elif cv == "block":
        bar(x0, y0 + int(ch * 0.55), x1, y1, primary)
        bar(x0, y0 + int(ch * 0.55), x1, y0 + int(ch * 0.55) + 3, accent)
        bar(x0 + 16, y0 + int(ch * 0.2), x1 - 24, y0 + int(ch * 0.2) + 9, dark)
        bar(x0 + 16, y0 + int(ch * 0.2) + 17, x0 + int(cw * 0.68), y0 + int(ch * 0.2) + 23, dark)
        bar(x0 + 16, y0 + int(ch * 0.45), x0 + int(cw * 0.58), y0 + int(ch * 0.45) + 5, muted)
    elif cv == "minimal":
        bar(x0 + 16, y0 + int(ch * 0.3), x0 + 50, y0 + int(ch * 0.3) + 4, accent)
        bar(x0 + 16, y0 + int(ch * 0.36), x1 - 24, y0 + int(ch * 0.36) + 9, dark)
        bar(x0 + 16, y0 + int(ch * 0.36) + 17, x0 + int(cw * 0.66), y0 + int(ch * 0.36) + 23, dark)
        bar(x0 + 16, y0 + int(ch * 0.6), x0 + int(cw * 0.52), y0 + int(ch * 0.6) + 5, muted)
    else:                                            # full
        bar(x0, y0, x1, y1, primary)
        bar(x0, y0 + int(ch * 0.78), x1, y0 + int(ch * 0.78) + 3, accent)
        bar(x0 + 16, y0 + int(ch * 0.4), x1 - 24, y0 + int(ch * 0.4) + 9, white)
        bar(x0 + 16, y0 + int(ch * 0.4) + 17, x0 + int(cw * 0.68), y0 + int(ch * 0.4) + 23, white)
        bar(x0 + 16, y0 + int(ch * 0.62), x0 + int(cw * 0.48), y0 + int(ch * 0.62) + 5, muted)

    # ---- 右：内容页骨架
    gx0, gy0, gx1, gy1 = x1 + gap, pad, x1 + gap + cw, pad + ch
    d.rounded_rectangle((gx0, gy0, gx1, gy1), radius=6, fill=page)
    chv = str(th.get("chrome", "band"))
    body_x = gx0 + 16
    if chv == "sidebar":
        sw = gx0 + int(cw * 0.09)
        bar(gx0, gy0, sw, gy1, primary)
        bar(sw, gy0, sw + 3, gy1, accent)
        body_x = gx0 + int(cw * 0.16)
        bar(body_x, gy0 + 18, body_x + int(cw * 0.5), gy0 + 26, dark)
        bar(body_x, gy0 + 32, body_x + int(cw * 0.28), gy0 + 37, muted)
    elif chv == "block":
        d.rounded_rectangle((gx0 + 16, gy0 + 14, gx0 + 16 + int(cw * 0.62), gy0 + 36),
                            radius=5, fill=primary)
        bar(gx0 + 24, gy0 + 22, gx0 + 24 + int(cw * 0.36), gy0 + 28, white)
        bar(gx0 + 16, gy0 + 44, gx0 + 16 + int(cw * 0.28), gy0 + 49, muted)
    elif chv == "minimal":
        bar(gx0 + 16, gy0 + 20, gx0 + 16 + int(cw * 0.55), gy0 + 29, dark)
        bar(gx0 + 16, gy0 + 34, gx0 + 16 + int(cw * 0.28), gy0 + 39, muted)
        bar(gx0 + 16, gy0 + 46, gx1 - 16, gy0 + 47, muted)
    else:                                            # band
        bar(gx0, gy0, gx1, gy0 + 4, accent)
        bar(gx0 + 16, gy0 + 20, gx0 + 16 + int(cw * 0.55), gy0 + 29, dark)
        bar(gx0 + 16, gy0 + 34, gx0 + 16 + int(cw * 0.28), gy0 + 39, muted)
        bar(gx0 + 16, gy0 + 44, gx0 + 46, gy0 + 46, accent)
    ty = gy0 + int(ch * 0.44)
    for i, frac in enumerate((0.78, 0.70, 0.8, 0.46)):
        bar(body_x, ty + i * 14, body_x + int(cw * frac), ty + i * 14 + 5, line)
    return img


def style_preview_png(key: str, out_path: Path, width: int = 420,
                      height: int = 250) -> Path:
    """预览图落盘版（GUI 弹窗与测试用）。"""
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    style_preview_image(key, width, height).save(out_path)
    return out_path
