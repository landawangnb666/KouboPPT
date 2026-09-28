"""学习笔记 / 题库 → Word（.docx），可选转 PDF。

公式与 PPT 同一策略：LaTeX 渲染成高清图片内嵌，Word 里绝不出现公式源码。
markdown 子集：#~### 标题、- 无序、1. 有序、$$..$$ 独立公式、| 表格 |。
"""
from __future__ import annotations

import re
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from . import formula, xmlsafe
from .subproc import NO_WINDOW

_MATH_LINE = re.compile(r"^\s*\$\$(.+?)\$\$\s*$")
_TABLE_SEP = re.compile(r"^\s*\|?[\s:-]+\|[\s:|-]*$")

Q_TYPE_LABEL = {"single": "单选题", "multi": "多选题", "blank": "填空题", "short": "解答题"}
Q_TYPE_ORDER = ("single", "multi", "blank", "short")

# 渲染失败的公式原文。Word 侧以前是静默降级：用户只看到纯文本，日志里一句提示都没有。
# 现在收集起来交给调用方打日志 + 写 `公式渲染失败.log`，与 PPT 侧行为一致。
# docx 由主线程串行生成（matplotlib 非线程安全），所以模块级列表足够。
_failed_formulas: list[str] = []


def reset_failed() -> None:
    """生成一份文档前调用，与 drain_failed() 配对。"""
    _failed_formulas.clear()


def drain_failed() -> list[str]:
    """取走并清空失败清单（生成完一份文档后调用）。"""
    got, _failed_formulas[:] = list(_failed_formulas), []
    return got


def _px_w_inches(png: Path) -> float:
    from PIL import Image
    with Image.open(png) as im:
        return im.size[0] / 300.0


def _setup_doc(doc: Document, accent: str = "1F4E79") -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    for i in (1, 2, 3):
        try:
            st = doc.styles[f"Heading {i}"]
            st.font.name = "Microsoft YaHei"
            st.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
            st.font.color.rgb = RGBColor.from_string(accent)
        except KeyError:
            pass


def _add_rich(par, text: str, size: int, formula_dir: Path) -> None:
    """向段落写入文字，$..$ 公式段渲染成内嵌图片。

    先清洗 XML 非法字符：模型偶尔会在题干/公式里混入控制字符，直接写进 Word
    会让 python-docx 抛 ValueError 并中止整轮任务。
    """
    text = xmlsafe.clean(text)
    for seg, is_math in formula.split_math(text):
        if not seg:
            continue
        if not is_math:
            run = par.add_run(seg)
            run.font.size = Pt(size)
            continue
        png = formula.render_dir_for(formula_dir, seg, size, "#2b2b2b", False)
        got = formula.render(seg, png, fontsize=size, color="#2b2b2b")
        if got is not None:
            run = par.add_run()
            run.add_picture(str(got), width=Inches(_px_w_inches(got)))
        else:
            # 渲染不了也不要甩源码：转成接近原文的纯文本（`a_{\text{Si}}` → `a_Si`）
            _failed_formulas.append(seg)
            run = par.add_run(formula.to_readable(seg))
            run.font.size = Pt(size)


def _add_math_paragraph(doc: Document, latex: str, formula_dir: Path, size: int = 12):
    latex = xmlsafe.clean(latex)
    png = formula.render_dir_for(formula_dir, latex, size, "#2b2b2b", True)
    got = formula.render(latex, png, fontsize=size, color="#2b2b2b", display=True)
    if got is None:
        _failed_formulas.append(latex)
        doc.add_paragraph(formula.to_readable(latex))
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(got), width=Inches(min(_px_w_inches(got), 6.0)))


# ---------------------------------------------------------------------------
def markdown_to_docx(md: str, out_path: Path, book_title: str = "",
                     chapter: str = "", formula_dir: Path | None = None,
                     title: str = "") -> Path:
    """Markdown → Word。title 直接指定文档标题（课后习题等非"学习笔记"文档用）。"""
    book_title = xmlsafe.clean(book_title)
    chapter = xmlsafe.clean(chapter)
    title = xmlsafe.clean(title)
    out_path = Path(out_path)
    formula_dir = Path(formula_dir) if formula_dir else out_path.parent / ".eq_cache"
    doc = Document()
    _setup_doc(doc)
    head = title or (f"{chapter} 学习笔记" if chapter else book_title)
    if head:
        doc.add_heading(head, level=0)
        if book_title and head != book_title:
            p = doc.add_paragraph(book_title)
            p.runs[0].font.color.rgb = RGBColor.from_string("8A99A8")

    lines = md.splitlines()
    i = 0
    first_block = True
    while i < len(lines):
        raw = lines[i].rstrip()
        i += 1
        if not raw.strip():
            continue
        m = _MATH_LINE.match(raw)
        if m:
            _add_math_paragraph(doc, m.group(1), formula_dir)
            first_block = False
            continue
        if raw.lstrip().startswith("|") and "|" in raw.lstrip()[1:]:
            rows = [raw]
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                rows.append(lines[i].rstrip())
                i += 1
            _add_table(doc, rows, formula_dir)
            first_block = False
            continue
        hm = re.match(r"^(#{1,4})\s+(.*)$", raw)
        if hm:
            # 开头的一级标题与文档标题重复，跳过
            if not (first_block and len(hm.group(1)) == 1):
                doc.add_heading(xmlsafe.clean(hm.group(2)), level=len(hm.group(1)))
            first_block = False
            continue
        p = doc.add_paragraph(style="List Bullet" if re.match(r"^\s*[-*]\s+", raw)
                              else "List Number" if re.match(r"^\s*\d+[.、]\s*", raw)
                              else None)
        text = re.sub(r"^\s*(?:[-*]|\d+[.、])\s*", "", raw)
        _add_rich(p, text, 11, formula_dir)
        first_block = False
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    return out_path


def _add_table(doc: Document, rows: list[str], formula_dir: Path) -> None:
    cells_rows = []
    for r in rows:
        if _TABLE_SEP.match(r):
            continue
        cells_rows.append([c.strip() for c in r.strip().strip("|").split("|")])
    if not cells_rows:
        return
    ncols = max(len(r) for r in cells_rows)
    table = doc.add_table(rows=len(cells_rows), cols=ncols)
    table.style = "Table Grid"
    for ri, cells in enumerate(cells_rows):
        for ci in range(ncols):
            text = cells[ci] if ci < len(cells) else ""
            _add_rich(table.cell(ri, ci).paragraphs[0], text, 10, formula_dir)
    doc.add_paragraph()


# ---------------------------------------------------------------------------
def questions_to_docx(questions: list[dict], out_path: Path, book_title: str = "",
                      chapter: str = "", with_answers: bool = True,
                      formula_dir: Path | None = None) -> Path:
    book_title = xmlsafe.clean(book_title)
    chapter = xmlsafe.clean(chapter)
    out_path = Path(out_path)
    formula_dir = Path(formula_dir) if formula_dir else out_path.parent / ".eq_cache"
    doc = Document()
    _setup_doc(doc)
    doc.add_heading(f"{chapter} 课后题库" if chapter else "题库", level=0)
    if book_title:
        doc.add_paragraph(book_title)

    groups: dict[str, list[dict]] = {t: [] for t in Q_TYPE_ORDER}
    for q in questions:
        groups.setdefault(str(q.get("type", "short")), []).append(q)

    for t in Q_TYPE_ORDER:
        qs = groups.get(t) or []
        if not qs:
            continue
        doc.add_heading(f"{Q_TYPE_LABEL.get(t, t)}", level=1)
        for n, q in enumerate(qs, 1):
            p = doc.add_paragraph()
            _add_rich(p, f"{n}. {q.get('stem', '').strip()}", 11, formula_dir)
            for opt in q.get("options") or []:
                op = doc.add_paragraph(style="List Bullet")
                _add_rich(op, str(opt), 11, formula_dir)

    if with_answers:
        doc.add_page_break()
        doc.add_heading("参考答案与解析", level=1)
        idx = 0
        for t in Q_TYPE_ORDER:
            for q in groups.get(t) or []:
                idx += 1
                p = doc.add_paragraph()
                _add_rich(p, f"{idx}. 答案：{q.get('answer', '略')}", 11, formula_dir)
                exp = str(q.get("explanation") or "").strip()
                if exp:
                    ep = doc.add_paragraph()
                    _add_rich(ep, f"   解析：{exp}", 10.5, formula_dir)
                    ep.runs[0].font.color.rgb = RGBColor.from_string("666666")

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    return out_path


# ---------------------------------------------------------------------------
_PS = r"""
param([string]$DocPath, [string]$PdfPath)
$ErrorActionPreference = 'Stop'
$word = New-Object -ComObject Word.Application
$word.Visible = $false
try {
    $doc = $word.Documents.Open($DocPath, $false, $true)
    $doc.SaveAs2($PdfPath, 17)          # wdFormatPDF
    $doc.Close($false)
} finally {
    $word.Quit()
    [System.Runtime.InteropServices.Marshal]::ReleaseComObject($word) | Out-Null
}
"""


def docx_to_pdf(docx_path: Path, pdf_path: Path | None = None) -> Path:
    """Word COM 转 PDF；本机没有 Word 时抛 RuntimeError（WPS 不支持该 ProgID 时同理）。"""
    docx_path = Path(docx_path)
    pdf_path = Path(pdf_path) if pdf_path else docx_path.with_suffix(".pdf")
    ps1 = docx_path.parent / "_docx2pdf.ps1"
    ps1.write_text(_PS, encoding="utf-8-sig")
    try:
        proc = subprocess.run(
            ["powershell", "-NoProfile", "-NonInteractive", "-ExecutionPolicy", "Bypass",
             "-File", str(ps1), "-DocPath", str(docx_path), "-PdfPath", str(pdf_path)],
            capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=180,
            **NO_WINDOW)
    finally:
        ps1.unlink(missing_ok=True)
    if not pdf_path.exists():
        raise RuntimeError("导出 PDF 失败（需要本机安装 Microsoft Word）。"
                           f"\n{proc.stderr[-400:] if proc.stderr else ''}")
    return pdf_path
