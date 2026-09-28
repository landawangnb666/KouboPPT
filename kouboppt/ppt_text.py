"""从 pptx 提取每页口播文稿。"""
from __future__ import annotations

import re

from pptx import Presentation

#: 支持的文稿来源
TEXT_SOURCES = ("正文", "备注", "备注优先")

_BULLET = re.compile(r"^[\s•·▪◦‣⁃\-–—*>]+")
_MULTI_PUNCT = re.compile(r"([，。！？；,.!?;:])\1+")


def load(path) -> Presentation:
    return Presentation(str(path))


def slide_count(prs: Presentation) -> int:
    return len(prs.slides._sldIdLst)


def slide_size_emu(prs: Presentation) -> tuple[int, int]:
    return int(prs.slide_width), int(prs.slide_height)


def _shape_texts(shape) -> list[str]:
    """递归收集一个形状里的文字（含组合形状和表格）。"""
    texts: list[str] = []
    if shape.shape_type == 6:  # GROUP
        for sub in shape.shapes:
            texts.extend(_shape_texts(sub))
        return texts
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            cells = [c.text.strip() for c in row.cells]
            line = "，".join(t for t in cells if t)
            if line:
                texts.append(line)
        return texts
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        lines = []
        for p in shape.text_frame.paragraphs:
            text = ("".join(r.text for r in p.runs) or p.text).strip()
            if text:
                lines.append(text)
        if lines:
            texts.append("\n".join(lines))
    return texts


def _body(slide) -> str:
    title_shape = None
    try:
        title_shape = slide.shapes.title
    except Exception:
        title_shape = None

    parts: list[str] = []
    if title_shape is not None:
        t = title_shape.text.strip()
        if t:
            parts.append(t)

    others = []
    for shape in slide.shapes:
        if title_shape is not None and shape.shape_id == title_shape.shape_id:
            continue
        others.append(shape)
    # 按版面位置从上到下、从左到右读
    others.sort(key=lambda s: (getattr(s, "top", 0) or 0, getattr(s, "left", 0) or 0))
    for shape in others:
        parts.extend(_shape_texts(shape))
    return "\n".join(parts)


def _notes(slide) -> str:
    try:
        if not slide.has_notes_slide:
            return ""
        tf = slide.notes_slide.notes_text_frame
        return tf.text.strip() if tf is not None else ""
    except Exception:
        return ""


def slide_script(slide, source: str) -> str:
    """按来源提取一页的口播文稿。"""
    if source not in TEXT_SOURCES:
        raise ValueError(f"未知文稿来源：{source}")
    if source == "备注":
        return _notes(slide)
    body = _body(slide)
    if source == "备注优先":
        notes = _notes(slide)
        return notes if notes.strip() else body
    return body


def clean_for_tts(text: str) -> str:
    """清理成适合朗读的一段话：去项目符号、合并行为自然句。"""
    lines: list[str] = []
    for raw in text.splitlines():
        line = _BULLET.sub("", raw.strip())
        if not line:
            continue
        line = re.sub(r"\s+", " ", line)
        if line:
            lines.append(line)
    if not lines:
        return ""
    joined = "，".join(lines)
    joined = re.sub(r"([。！？；!?;])，", r"\1", joined)  # 句号后不再加逗号
    joined = _MULTI_PUNCT.sub(r"\1", joined)              # 连续重复标点去重
    return joined.strip("，, ")


def extract_all(prs: Presentation, source: str) -> list[str]:
    """返回每页的待朗读文本（未 clean）。"""
    return [slide_script(slide, source) for slide in prs.slides]
