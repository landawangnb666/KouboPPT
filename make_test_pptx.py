"""生成测试样例 PPT（16:9，6 页：含正文、备注、空页）。"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Inches, Pt

W, H = Emu(12192000), Emu(6858000)  # 16:9


def add_slide(prs, title, bullets, notes=None, blank=False):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # 空白版式
    if not blank:
        tb = slide.shapes.add_textbox(Inches(0.6), Inches(0.5), Inches(8.8), Inches(1.2))
        p = tb.text_frame.paragraphs[0]
        p.text = title
        p.font.size = Pt(36)
        p.font.bold = True
        p.font.color.rgb = RGBColor(0x1F, 0x3B, 0x73)
        body = slide.shapes.add_textbox(Inches(0.9), Inches(2.0), Inches(8.2), Inches(4.2))
        tf = body.text_frame
        tf.word_wrap = True
        for i, b in enumerate(bullets):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.text = "• " + b
            para.font.size = Pt(22)
            para.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
            para.space_after = Pt(14)
    if notes:
        slide.notes_slide.notes_text_frame.text = notes
    return slide


def main():
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H

    add_slide(prs, "口播PPT 测试样例", ["这是一个自动生成的测试演示文稿", "用于验证 PPT 转口播视频的完整流程"],
              notes="开场白：大家好，这是一段写在备注里的讲稿，用来验证备注来源。")
    add_slide(prs, "第一部分：项目背景",
              ["传统课程视频制作需要录音、剪辑、合成多个步骤",
               "本工具把这三步自动化为一次点击",
               "支持批量处理整个课程目录"])
    add_slide(prs, "第二部分：技术方案",
              ["幻灯片渲染：调用本机 PowerPoint 导出高清图片",
               "语音合成：内置 Edge TTS，也可扩展其他引擎",
               "视频合成：ffmpeg 逐页合成后无损拼接"])
    add_slide(prs, "", [], blank=True)  # 无文字页 → 应停留数秒
    add_slide(prs, "第三部分：使用方法",
              ["选择一个或多个 PPT 文件",
               "填写页码范围可以拆分出多节课视频",
               "点击开始生成，等待完成即可"])
    add_slide(prs, "谢谢观看", ["祝大家使用愉快"])

    out = Path(__file__).parent / "test_out" / "样例课程_第1课.pptx"
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    print("已生成:", out)


if __name__ == "__main__":
    main()
