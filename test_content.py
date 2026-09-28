"""离线冒烟测试：不依赖 AI Key，验证公式渲染/PPT 构建/Word 导出/教材切分。"""
import json
import shutil
import sys
from pathlib import Path

from kouboppt import formula, pipeline, video

# 控制台是 GBK 时打印 ✔/⚠ 会 UnicodeEncodeError，导致测试整个失败
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(errors="replace")

OUT = Path("test_out")
OUT.mkdir(exist_ok=True)

r1 = formula.render(r"\int_0^\infty e^{-x^2}dx = \frac{\sqrt{\pi}}{2}",
                    OUT / "eq1.png", display=True)
r2 = formula.render("其中 $E=mc^2$ 为质能方程", OUT / "eq2.png")
r3 = formula.render(r"\begin{cases} x+y=1 \\ x-y=3 \end{cases}",
                    OUT / "eq3.png", display=True)
r4 = formula.render(r"\begin{pmatrix} 1 & 2 \\ 3 & 4 \end{pmatrix}",
                    OUT / "eq4.png", display=True)
r6 = formula.render(r"A = \begin{bmatrix} a_{11} & a_{12} \\ a_{21} & a_{22} \end{bmatrix}",
                    OUT / "eq6.png", display=True)
r7 = formula.render(r"f(x) = \begin{cases} x, & x \geq 0 \\ -x, & x < 0 \end{cases}",
                    OUT / "eq7.png", display=True)
r5 = formula.render(r"\qqqbadcmd{x}", OUT / "eq5.png", display=True)  # 应为 None

print("积分:", r1)
print("混排:", r2)
print("cases:", r3)
print("pmatrix:", r4)
print("bmatrix带前缀:", r6)
print("cases带等式:", r7)
print("坏公式(应None):", r5)
assert all((r1, r2, r3, r4, r6, r7)), "有公式渲染失败"
assert r5 is None, "坏公式应返回 None"

# ---------------- 公式归一化：模型/OCR 常见写法必须能渲染 ----------------
# mathtext 不认 \le/\ge、\textbf、\frac1n 省花括号、空单元格、gathered、多行 $$，
# 由 sanitize_latex 归一后再渲染。这里锁住这些曾经失败的类型，防止回归。
_eq_ok = [
    r"$x\le 0$", r"$a\le b\le c$", r"$x\ge 0$", r"$\land$", r"$\lor$", r"$\gets$",
    r"$a\bmod b$", r"$a\pmod{n}$", r"$\textbf{ab}$", r"$\textit{ab}$", r"$\textrm{ab}$",
    r"$\bm{v}$", r"$\underbrace{x}$", r"$\overbrace{x}$", r"$\tfrac{a}{b}$",
    r"$\big(x\big)$", r"$\Big[x\Big]$", r"$\lVert x\rVert$", r"$\stackrel{a}{=}$",
    r"$\xrightarrow{f}$", r"$\xleftarrow{g}$", r"$\displaystyle\sum$",
    r"$\sum\limits_{i=1}^n$", r"$\frac1n$", r"$\frac12$", r"$\sqrt2$",
    r"$\frac1{n^2}$", r"$\sum_{n=1}^{\infty}\frac1{n^2}=\frac{\pi^2}{6}$",
    r"$\left(1+\frac1n\right)^n$", r"$\forall x\in\mathbb{R},\ x^2\ge0$",
    r"$$\begin{gathered} a = b \\ c = d \end{gathered}$$",
    r"$$\begin{cases} \text{是} & x>0 \\ \text{否} & x\le 0 \end{cases}$$",
    r"$$\begin{matrix} a & \\ c & d \end{matrix}$$",
    r"$$\begin{align} a &= b \\ &= c \end{align}$$",
    "$$\nE=mc^2\n$$",                              # 真实换行的多行 $$
]
for _i, _tex in enumerate(_eq_ok):
    _got = formula.render(_tex, OUT / f"eqn{_i}.png", display=_tex.startswith("$$"))
    assert _got is not None, f"公式归一化后仍渲染失败：{_tex}"

# 归一化只按命令词边界替换，不误伤长命令；且幂等
assert formula.sanitize_latex(r"\left(x\right)") == r"\left(x\right)", "误伤 \\left"
assert formula.sanitize_latex(r"x\leq y") == r"x\leq y", "误伤 \\leq"
assert formula.sanitize_latex(r"\frac{a}{b}") == r"\frac{a}{b}", "误伤已带括号的分式"
assert formula.sanitize_latex(r"\textbf{ab}") == r"\mathbf{ab}"
assert formula.sanitize_latex(r"a\bmod b") == r"a\ \mathrm{mod} b", "误伤 \\bmod"
_once = formula.sanitize_latex(r"$\sqrt2\le\frac1n$")
assert formula.sanitize_latex(_once) == _once, f"归一化不幂等：{_once!r}"
# 文本命令：纯 ASCII 保留命令本体（保住粗斜体），中文提取到 math 之外
assert formula._math_body(r"\mathbf{v}") == r"$\mathbf{v}$"
assert formula._math_body(r"\mathrm{速度}") == "速度"
print("公式归一化 OK（le/ge、省花括号、textbf、空单元格、多行 $$、gathered）")

# ---------------- 含中文的公式：曾整行降级，现改用中文字体集渲染 ----------------
# `E_{\text{总}}`、`x_{中}` 这类"中文出现在公式组内"的写法，STIX 没有中文字形；
# 旧实现在括号不配平时直接返回 None → 整行降级成纯文本 `E_总`（物理/化学教材里极常见）。
from PIL import Image as _Image  # noqa: E402

_cjk_eq = [
    r"$E_{\text{总}}$", r"$x_{中}$", r"$v_{平均}=\frac{s}{t}$",
    r"$E_{\text{总}}=mc^2$", r"$\rho_{\text{水}}=1.0\times10^{3}\ \mathrm{kg/m^3}$",
    r"$$a_{\text{最大}}=\frac{F}{m}$$",
]
for _i, _tex in enumerate(_cjk_eq):
    _got = formula.render(_tex, OUT / f"cjk{_i}.png", display=_tex.startswith("$$"))
    assert _got is not None, f"含中文的公式仍渲染失败：{_tex}"
assert formula._CJK.search(formula._math_body(r"E_{\text{总}}") or ""), "中文应留在 math 里"
assert formula._cjk_font(), "应能挑到一个可用的中文字体"
# 不含中文的公式不该切字体集（保住 STIX 的数学斜体）
assert not formula._CJK.search(formula._math_body(r"E_{total}=mc^2") or "")

# 公式图片缓存 key 必须带 CACHE_VERSION：规则升级后旧图自动失效，
# 否则像 `$x^22$`（旧规则画成 x²2 却"成功"落盘）修好了也仍显示旧图。
_p1 = formula.render_dir_for(OUT, r"$x^2$", 20, "#000000", False)
_ver = formula.CACHE_VERSION
formula.CACHE_VERSION = _ver + 1
_p2 = formula.render_dir_for(OUT, r"$x^2$", 20, "#000000", False)
formula.CACHE_VERSION = _ver
assert _p1 != _p2, "CACHE_VERSION 没参与缓存 key"

# 定界符：字形按公式字号渲染后"切片拉伸"，笔画不会被一起拉粗；缓存落在 _delims/ 子目录
_small = OUT / "_brace_small.png"
formula._char_png("{", 24, "#000000", _small, base_px=20)
_si = _Image.open(_small)
_st = formula._stretch_v(_si, 900, nib=True)
assert abs(_st.height - 900) <= 2 and _st.width == _si.width, (_st.size, _si.size)
assert formula._stretch_v(_si, 10, nib=True).height == _si.height, "目标更矮应原样返回"
_env_cache = OUT / "_env_probe"          # 干净目录，免得 test_out 里历次遗留的临时文件干扰断言
shutil.rmtree(_env_cache, ignore_errors=True)
_brace_tex = r"$$\begin{cases} a=b \\ c=d \\ e=f \\ g=h \\ i=j \end{cases}$$"
assert formula.render(_brace_tex, _env_cache / "brace.png", fontsize=20,
                      display=True), "长 cases 渲染失败"
assert list((_env_cache / "_delims").glob("*.png")), "定界符应落在 _delims/ 子目录"
assert not list(_env_cache.glob("_delim_*.png")), "不该再直接在缓存根目录生成 _delim_*.png"
# 环境两侧的前缀/后缀图曾用固定文件名 prefix.png/suffix.png（会被互相覆盖），现在带内容哈希
_side = r"$$F = \begin{matrix} 1 & 2 \\ 3 & 4 \end{matrix}$$"
assert formula.render(_side, _env_cache / "side.png", fontsize=20,
                      display=True), "带前缀的矩阵渲染失败"
assert not (_env_cache / "_env_cells" / "prefix.png").exists(), "不该再生成固定名 prefix.png"
assert list((_env_cache / "_env_cells").glob("prefix_*.png")), "前缀图应带内容哈希"
shutil.rmtree(_env_cache, ignore_errors=True)
print("公式中文渲染 OK（中文下标、缓存版本号、定界符切片拉伸、前缀图哈希命名）")

# ---------------- 模型 JSON 宽容解析 + 公式定界符归一 ----------------
from kouboppt.llm import extract_json

bad = ('{"slides":[{"title":"生灭过程","lines":["只允许 \\(+1\\)（生）或 \\(-1\\)（灭）",'
       '"出生率 \\lambda_i 与 \\mu"],"script":"读作 二分之一"}]}')
j = extract_json(bad)                     # \( \l 非法转义 → 自动补齐
assert j["slides"][0]["title"] == "生灭过程", j
assert j["slides"][0]["lines"][0] == "只允许 \\(+1\\)（生）或 \\(-1\\)（灭）", j

j2 = extract_json(r'{"s": "已转义 \\(x\\) 与未转义 \(y\) 混排"}')   # 合法转义不能被误伤
assert j2["s"] == r"已转义 \(x\) 与未转义 \(y\) 混排", j2

j3 = extract_json('{"a":"第一行\n第二行","b":"制表\t符"}')          # 字符串内裸换行/制表符
assert j3["a"] == "第一行\n第二行" and j3["b"] == "制表\t符", j3

j4 = extract_json('```json\n{"ok": true, "s": "line\\nbreak"}\n```')
assert j4 == {"ok": True, "s": "line\nbreak"}, j4

assert formula.normalize_delims(r"设 \(x>0\)，独立公式 \[E=mc^2\] 结束") == \
    "设 $x>0$，独立公式 $$E=mc^2$$ 结束"
assert formula.has_math(r"设 \(x>0\)")
assert not formula.has_math("没有公式的一行")
print("模型 JSON 宽容解析 + 定界符归一 OK")

from kouboppt import slidegen

sample = slidegen.LessonSpec(
    title="第1章 测试课",
    subtitle="高等数学（上册）",
    sections=[
        slidegen.SectionSpec(
            title="函数与极限",
            slides=[
                slidegen.SlideSpec(
                    title="极限的定义",
                    lines=[
                        "设函数 f(x) 在 x0 邻域内有定义",
                        "$$\\lim_{x \\to x_0} f(x) = A$$",
                        "严格语言：对任意 $\\varepsilon > 0$，存在 $\\delta > 0$",
                    ],
                    script="同学们好，这节课我们学习极限的严格定义。",
                ),
            ],
        ),
    ],
)
pptx_path = slidegen.build(sample, OUT / "sample_lesson.pptx", theme="stem_blue")
print("PPT:", pptx_path, pptx_path.stat().st_size, "bytes")

# ---------------- PPT 风格骨架（配色 × 版式）----------------
from pptx import Presentation as _Prs

assert set(slidegen.style_labels()) == set(slidegen.THEMES), "风格库的每个 key 都要有中文名"
assert slidegen.DEFAULT_STYLE in slidegen.THEMES, "兜底风格必须在库里"
_chromes = {t.get("chrome") for t in slidegen.THEMES.values()}
_covers = {t.get("cover") for t in slidegen.THEMES.values()}
assert {"band", "sidebar", "block", "minimal"} <= _chromes, f"四套页眉骨架都要有：{_chromes}"
assert {"full", "split", "block", "minimal"} <= _covers, f"四套封面骨架都要有：{_covers}"

# 同一份内容换任何风格，页数都不能变（骨架只换皮不改内容/分页）
_style_lesson = slidegen.LessonSpec(
    title="第1章 测试课", subtitle="测试教材",
    sections=[
        slidegen.SectionSpec(title="第一节", slides=[
            slidegen.SlideSpec(title="概念", lines=["要点一", "$$a^2+b^2=c^2$$"],
                               script="稿一"),
            slidegen.SlideSpec(title="例题", lines=["要点二"], script="稿二"),
        ]),
        slidegen.SectionSpec(title="第二节", slides=[
            slidegen.SlideSpec(title="小结", lines=["要点三"], script="稿三"),
        ]),
    ])
_style_pages = {}
for _k in slidegen.THEMES:
    _p = slidegen.build(_style_lesson, OUT / f"style_{_k}.pptx", theme=_k)
    _style_pages[_k] = len(_Prs(str(_p)).slides._sldIdLst)
assert len(set(_style_pages.values())) == 1, f"各风格页数应一致：{_style_pages}"

# 预览图：离线可画（GUI 候选弹窗用），未知风格兜底出图不报错
_pv = slidegen.style_preview_png("medical_teal", OUT / "style_preview.png")
assert _pv.exists() and _pv.stat().st_size > 500, _pv
assert slidegen.style_preview_image("no_such_style").size == (420, 250), "未知风格应兜底"
print("PPT 风格骨架 OK（%d 套：配色 × 版式，页数一致，预览图可离线生成）"
      % len(slidegen.THEMES))

# ---------------- 自动选风格：AI 挑 + 落盘缓存 + 失败兜底 ----------------
class _StyleLLM:
    """只应答复"选风格"的桩；n 用来验证缓存命中后不再问 AI。"""

    def __init__(self, picks):
        self.picks, self.n = picks, 0

    def chat_json(self, system, user, images=None, max_tokens=0, **kw):
        self.n += 1
        assert "可选风格" in user, f"选风格请求该带上风格清单：{user[:120]}"
        return {"styles": self.picks}


_style_book = OUT / "_style_cache_book"
_bad_book = OUT / "_style_cache_book_dead"
shutil.rmtree(_style_book, ignore_errors=True)       # 上一次跑留下的缓存会干扰"第一次问 AI"
shutil.rmtree(_bad_book, ignore_errors=True)
_style_llm = _StyleLLM([{"key": "medical_teal", "reason": "医学类教材"},
                        {"key": "不存在的风格", "reason": "非法项该被过滤"},
                        {"key": "medical_teal", "reason": "重复项该被去掉"}])
assert slidegen.resolve_style(_style_llm, "人体解剖学", _style_book,
                              log=lambda s: None) == "medical_teal"
assert _style_llm.n == 1, "第一次该问 AI"
assert slidegen.resolve_style(_style_llm, "人体解剖学", _style_book,
                              log=lambda s: None) == "medical_teal"
assert _style_llm.n == 1, "第二次该命中 PPT风格.json 缓存，不再问 AI"
assert (slidegen.load_style_choice(_style_book) or {}).get("key") == "medical_teal"


class _DeadStyleLLM:
    def chat_json(self, *a, **kw):
        raise RuntimeError("接口挂了")


assert slidegen.resolve_style(_DeadStyleLLM(), "某本怪书", _bad_book,
                              log=lambda s: None) == slidegen.DEFAULT_STYLE
assert slidegen.load_style_choice(_bad_book) is None, "失败不该写缓存"
shutil.rmtree(_style_book, ignore_errors=True)       # 测试目录不留缓存，回头重跑仍从"第一次问 AI"开始
print("自动选风格 OK（AI 挑 / 缓存命中只问一次 / 失败回落默认风格）")

from kouboppt import docgen

md = """# 第1章 学习笔记

## 1.1 极限
- 定义：$\\lim_{x \\to x_0} f(x) = A$
- 重要极限：$$\\lim_{x \\to 0} \\frac{\\sin x}{x} = 1$$

## 1.2 练习题提示
背住两个重要极限。
"""
questions = [
    {"type": "single", "stem": "极限 $\\lim_{x\\to 0}\\sin(1/x)$ 是？",
     "options": ["A. 0", "B. 1", "C. 不存在", "D. $\\infty$"],
     "answer": "C", "explanation": "震荡无极限。"},
    {"type": "blank", "stem": "$\\frac{d}{dx}x^n = $ ____",
     "answer": "$nx^{n-1}$", "explanation": "幂函数求导。"},
]
note_path = docgen.markdown_to_docx(md, OUT / "sample_notes.docx",
                                     book_title="测试教材",
                                     chapter="第1章", formula_dir=OUT / "eqcache")
quiz_path = docgen.questions_to_docx(questions, OUT / "sample_quiz.docx",
                                      book_title="测试教材", chapter="第1章",
                                      with_answers=True, formula_dir=OUT / "eqcache")
print("笔记:", note_path, note_path.stat().st_size)
print("题库:", quiz_path, quiz_path.stat().st_size)

# ---------------- 教材切分纯函数（不依赖 AI / 网络） ----------------
from kouboppt import textbook as tb

# 按每节页数：300 页 ÷ 15 = 20 节；307 页的余数独立成最后一节
ls = tb.split_by_pages(300, 1, 15)
assert len(ls) == 20 and (ls[0].start, ls[0].end) == (1, 15), ls
assert ls[0].title == "第1课 p1-p15", ls[0].title
assert (ls[-1].start, ls[-1].end) == (286, 300), ls[-1]
ls = tb.split_by_pages(307, 1, 15)
assert len(ls) == 21 and (ls[-1].start, ls[-1].end) == (301, 307), ls[-1]
ls = tb.split_by_pages(40, 21, 15)
assert [(c.start, c.end) for c in ls] == [(21, 35), (36, 40)], ls
try:
    tb.split_by_pages(5, 6, 15)
    raise AssertionError("起始页超出总页数应报错")
except ValueError:
    pass
chs = [tb.Chapter("第1章 绪论", 1, 20), tb.Chapter("第2章 极限", 21, 40)]
ls = tb.split_by_pages(40, 1, 15, chapters=chs)
assert "第1章 绪论" in ls[0].title and "第2章 极限" in ls[2].title, [c.title for c in ls]

# 按教材章节：章内细分、不跨章、小章独立成节
ls = tb.split_chapters([tb.Chapter("第1章", 1, 20), tb.Chapter("第2章", 21, 22)], 15)
assert [(c.title, c.start, c.end) for c in ls] == [
    ("第1章 第1节", 1, 15), ("第1章 第2节", 16, 20), ("第2章", 21, 22)], ls

# 章节结构.json 读写
tmp = OUT / "_chapters_test"
if tmp.exists():
    shutil.rmtree(tmp)
tmp.mkdir()
tb.save_chapter_structure(tmp, chs, "ai", 40)
got, src = tb.load_chapter_structure(tmp, 40)
assert src == "ai" and [(c.title, c.start, c.end) for c in got] == [
    ("第1章 绪论", 1, 20), ("第2章 极限", 21, 40)], got
tb.chapter_structure_path(tmp).unlink()
assert tb.load_chapter_structure(tmp, 40) is None
tb.chapter_structure_path(tmp).write_text("{坏 JSON", encoding="utf-8")
assert tb.load_chapter_structure(tmp, 40) is None       # 坏文件被忽略、不抛错
tb.chapter_structure_path(tmp).write_text(json.dumps({"chapters": [
    {"title": "b", "start": 15, "end": 20}, {"title": "a", "start": 1, "end": 10}]},
    ensure_ascii=False), encoding="utf-8")
got, _ = tb.load_chapter_structure(tmp, 25)             # 乱序→排序；缝隙并入前章；末章补全
assert [(c.title, c.start, c.end) for c in got] == [("a", 1, 14), ("b", 15, 25)], got
shutil.rmtree(tmp)
print("切分纯函数 OK（余数、起始页、章内细分、章节结构.json）")

# ---------------- 小节对齐 + 短节合并 ----------------
# AI 目录解析 + 归一化：缺起点/倒置/越界条目丢弃，重叠小节后移让不下则丢弃
parsed = tb.normalize_chapters(tb._parse_toc_chapters({"chapters": [
    {"title": "第1章", "start": 3, "end": 40,
     "sections": [{"title": "1.1 a", "start": 3, "end": 10},
                  {"title": "无起点", "end": 11},
                  {"title": "倒置", "start": 15, "end": 12},
                  {"title": "越界", "start": 999, "end": 1200},
                  {"title": "1.2 b", "start": 12, "end": 20},
                  {"title": "重叠", "start": 13, "end": 15}]}]}, 100), 100)
assert (parsed[0].start, parsed[0].end) == (3, 100)   # 末章补到全书末尾
assert [(s.title, s.start, s.end) for s in parsed[0].sections] == [
    ("1.1 a", 3, 10), ("1.2 b", 12, 20)], parsed[0].sections

# 有小节的章：沿小节边界聚合，每节 ≤ per 页，小节不被拦腰切断
# （1.1+1.2=18≤20 一节；再加 1.3 就 27>20 → 在小节边界断开；1.3+1.4=18 一节）
ch = tb.Chapter("第1章", 1, 45, [
    tb.Chapter("1.1 a", 1, 9), tb.Chapter("1.2 b", 10, 18),
    tb.Chapter("1.3 c", 19, 27), tb.Chapter("1.4 d", 28, 36),
    tb.Chapter("1.5 e", 37, 45)])
ls = tb.split_chapters([ch], 20)
assert [(c.title, c.start, c.end) for c in ls] == [
    ("1.1 a", 1, 18), ("1.3 c", 19, 36), ("1.5 e", 37, 45)], ls

# 单小节超过 per 页：独占一节但均分成多块（标题加 i/n），后小节不受牵连
ch2 = tb.Chapter("第2章", 46, 100, [tb.Chapter("2.1 大", 46, 90),
                                    tb.Chapter("2.2 末", 91, 100)])
ls = tb.split_chapters([ch2], 20)
assert [(c.title, c.start, c.end) for c in ls] == [
    ("2.1 大（1/3）", 46, 60), ("2.1 大（2/3）", 61, 75), ("2.1 大（3/3）", 76, 90),
    ("2.2 末", 91, 100)], ls

# 短节合并：尾节 5 页 → 并入上一节；阈值更大时逐步并成整章
ls = tb.split_chapters([tb.Chapter("第1章", 1, 45)], 20, min_pages=12)
assert [(c.title, c.start, c.end) for c in ls] == [
    ("第1章 第1节", 1, 20), ("第1章 第2节", 21, 45)], ls
ls = tb.split_chapters([tb.Chapter("第1章", 1, 45)], 20, min_pages=25)
assert [(c.title, c.start, c.end) for c in ls] == [("第1章", 1, 45)], ls

# 短章合并：首章一节太短 → 并入下一章；末章太短 → 并入上一章
ls = tb.split_chapters([tb.Chapter("引言", 1, 5), tb.Chapter("主体", 6, 44),
                        tb.Chapter("尾声", 45, 48)], 20, min_pages=8)
assert [(c.title, c.start, c.end) for c in ls] == [
    ("主体 第1节", 1, 25), ("主体 第2节", 26, 48)], ls

# 章节结构.json v2：带 sections 往返一致（幂等）
tmp = OUT / "_chapters_sec_test"
if tmp.exists():
    shutil.rmtree(tmp)
tmp.mkdir()
tb.save_chapter_structure(tmp, [ch], "ai", 45)
got, src = tb.load_chapter_structure(tmp, 45)
assert src == "ai" and got == [ch], got
shutil.rmtree(tmp)

# PDF 二级书签 → 章内小节
import pymupdf as fitz
bdoc = fitz.open()
for _ in range(4):
    bdoc.new_page()
bdoc.set_toc([[1, "第1章", 1], [2, "1.1 a", 1], [2, "1.2 b", 3],
              [1, "第2章", 4], [2, "2.1 c", 4]])
bchs = tb.chapters_from_bookmarks(bdoc)
assert [(c.title, c.start, c.end) for c in bchs] == [
    ("第1章", 1, 3), ("第2章", 4, 4)], bchs
assert [(s.title, s.start, s.end) for s in bchs[0].sections] == [
    ("1.1 a", 1, 2), ("1.2 b", 3, 3)], bchs[0].sections
assert [(s.title, s.start, s.end) for s in bchs[1].sections] == [
    ("2.1 c", 4, 4)], bchs[1].sections
bdoc.close()
print("小节对齐 + 短节合并 OK（AI 解析容错、小节聚合、超长切块、短节/短章合并、缓存 v2、二级书签）")

# ---------------- 章节归一化幂等：首次识别 === 重跑读缓存 边界必须一致 ----------------
# 旧实现：书签路径留缝隙、load 按不同方向归一化，导致重跑边界漂移、目录名错位、
# 末章被撑到全书末尾、甚至目录名相同页码却变了（静默复用缺页旧稿）。
# 现在三条来源共用 normalize_chapters，save→load 必须往返一致。
norm_dir = OUT / "_norm_test"
if norm_dir.exists():
    shutil.rmtree(norm_dir)
norm_dir.mkdir()

def _save_load(raw, total):
    tb.save_chapter_structure(norm_dir, raw, "bookmarks", total)
    out, _ = tb.load_chapter_structure(norm_dir, total)
    return [(c.title, c.start, c.end) for c in out]

# 直接校验纯函数
assert tb.normalize_chapters(
    [tb.Chapter("第1章 A", 1, 30), tb.Chapter("第2章 B", 32, 60)], 80) == [
    tb.Chapter("第1章 A", 1, 31), tb.Chapter("第2章 B", 32, 80)], "normalize 结果不符"

# 场景 A：中间有 1 页碎章被丢弃留缝隙；save→load 再读一次必须完全一致
sca = [tb.Chapter("第1章 A", 1, 30), tb.Chapter("第2章 B", 32, 60)]
r1 = _save_load(sca, 80)
r2, _ = tb.load_chapter_structure(norm_dir, 80)        # 用存好的 JSON 再读一遍
r2_t = [(c.title, c.start, c.end) for c in r2]
assert r1 == r2_t, (r1, r2_t)
assert r1 == [("第1章 A", 1, 31), ("第2章 B", 32, 80)], r1
# 切课首次与重跑目录名/页码都必须一致（不再因错位触发整章重识别）
r1c = [tb.Chapter(t, a, b) for (t, a, b) in r1]
assert [(c.title, c.start, c.end) for c in tb.split_chapters(r1c, 15)] == \
       [(c.title, c.start, c.end) for c in tb.split_chapters(r2, 15)]

# 场景 B：缝隙并入前一章后目录名不变但页码变了 —— 必须首次=重跑（否则会静默复用旧稿）
scb = [tb.Chapter("第1章 A", 1, 10), tb.Chapter("第2章 B", 12, 30)]
assert _save_load(scb, 30) == [("第1章 A", 1, 11), ("第2章 B", 12, 30)]

# 手工编辑留缝隙：统一并入前一章，与 AI 路径语义一致
scc = [tb.Chapter("第1章", 1, 10), tb.Chapter("第3章", 40, 50)]
assert tb.normalize_chapters(scc, 50) == [
    tb.Chapter("第1章", 1, 39), tb.Chapter("第3章", 40, 50)], "手工缝隙方向不符"
shutil.rmtree(norm_dir)
print("章节归一化幂等 OK（save→load 往返一致，首次=重跑，缝隙统一并入前章）")

# ---------------- 课后习题切分纯函数 ----------------
md_lines = [
    "# 第1章 绪论",
    "直角三角形的两条直角边的平方和等于斜边的平方。",
    "### 本章小结",
    "勾股定理是几何学的基石之一。",
    "## 课后习题",                                    # ← 习题区（第 4 行）开始
    "1. 已知两直角边为 3 和 4，求斜边的长度。",
    "2. 证明勾股定理在任意直角三角形中成立。",
    "## 思考题",
    "3. 为什么该定理在古代中国被称为商高定理？",
    "参考答案：略。",                                 # ← 习题区结束（第 9 行）
    "# 第2章 数列",
    "数列是定义在正整数集上的函数。",
]
md = "\n".join(md_lines)
body, ex = tb.split_exercises(md)
assert ex == "\n".join(md_lines[4:10]), ex
assert body == "\n".join(md_lines[:4] + md_lines[10:]), body
ex_sample = ex

# 没有习题：原样返回、习题为空
body, ex = tb.split_exercises("\n".join(md_lines[:4]))
assert ex == "" and body == "\n".join(md_lines[:4]), (body, ex)

# 关键词只在正文里顺带出现（长句/长行）不算标题：正文含"作业"二字也不切
body, ex = tb.split_exercises("# 标题\n这一节我们先讲清楚习题课该怎么上，再说作业怎么布置。\n")
assert ex == "" and "习题课" in body, (body, ex)

# 关键词标题但内容太短（<30 字）视为误判，不切
body, ex = tb.split_exercises("# 标题\n正文正文正文。\n## 思考题\n略。\n")
assert ex == "" and "## 思考题" in body, (body, ex)

# 习题区到下一个标题行为止（##/### 都算边界）
body, ex = tb.split_exercises("# 标题\n正文。\n### 复习题\n1. 什么是导数？这是一个足够长的题目描述文字。\n"
                              "## 下一节\n下一节正文继续讲。\n")
assert ex.startswith("### 复习题") and "下一节正文" in body, (body, ex)
print("课后习题切分 OK（标题行判定、习题区边界、误判保护）")

# 课后习题单独出 Word（docgen 直接指定标题，不再叫"学习笔记"）
ex_doc = docgen.markdown_to_docx(ex_sample, OUT / "sample_exercises.docx",
                                 book_title="测试教材", formula_dir=OUT / "eqcache",
                                 title="第1课 p1-p15 课后习题")
from docx import Document as _Doc

paras = [p.text for p in _Doc(str(ex_doc)).paragraphs]
assert paras[0] == "第1课 p1-p15 课后习题", paras[:3]
assert paras[1] == "测试教材", paras[:3]
assert "商高定理" in "\n".join(paras), paras
note_paras = [p.text for p in _Doc(str(note_path)).paragraphs]
assert note_paras[0] == "第1章 学习笔记" and note_paras[1] == "测试教材", note_paras[:3]
print("课后习题文档 OK（标题就位、题目原文在内；学习笔记标题照旧）")

# ------------------------------------------------- 视频帧率：默认 10fps，可覆盖
# 静态幻灯片 10fps 比 30fps 快约 5.9 倍且画质更高（实测见 test_out/_bench_enc7.py）
assert pipeline.Settings().fps == 10, pipeline.Settings().fps
assert video.make_segment.__defaults__[0] == 10, video.make_segment.__defaults__

captured: list[list[str]] = []


class _FakeProc:
    returncode = 0
    stderr = ""


def _fake_run(cmd, **kw):
    captured.append(cmd)
    Path(cmd[-1]).write_bytes(b"fake mp4")      # make_segment 会检查输出文件存在
    return _FakeProc()


_orig_run = video.subprocess.run
video.subprocess.run = _fake_run
try:
    seg = OUT / "_seg_fps.mp4"
    video.make_segment(OUT / "eq1.png", None, 3.0, seg, 1920, 1080)
    assert captured[-1][captured[-1].index("-framerate") + 1] == "10", captured[-1]
    video.make_segment(OUT / "eq1.png", None, 3.0, seg, 1920, 1080, fps=30)
    assert captured[-1][captured[-1].index("-framerate") + 1] == "30", captured[-1]
    seg.unlink(missing_ok=True)
finally:
    video.subprocess.run = _orig_run
print("视频帧率 OK（默认 10fps，Settings/调用方可覆盖）")

# ------------------------------------------------- 编码器选择、参数与硬编失败回退
for codec, fmt in (("libx264", "yuv420p"), ("h264_nvenc", "yuv420p"),
                   ("h264_qsv", "nv12"), ("h264_amf", "nv12")):
    assert video.codec_pix_fmt(codec) == fmt, codec
for codec, opt in (("libx264", "-crf"), ("h264_nvenc", "-cq"),
                   ("h264_qsv", "-global_quality"), ("h264_amf", "-qp_i")):
    args = video.codec_args(codec)
    assert args[args.index("-c:v") + 1] == codec, args
    assert opt in args, (codec, args)
assert "stillimage" in video.codec_args("libx264")
assert "-crf" not in video.codec_args("h264_qsv"), "硬编不认 -crf"

video.set_selfcheck({"timings": {"libx264": 1.1, "h264_qsv": 2.0},
                     "hw_available": ["h264_qsv"], "recommended": "libx264"})
assert video.resolve_codec("auto") == "libx264"
assert video.resolve_codec("hw") == "h264_qsv"
assert video.resolve_codec("sw") == "libx264"
assert video.resolve_codec("h264_nvenc") == "h264_nvenc", "具体编码器名应直接透传"
video.set_selfcheck({"timings": {"h264_nvenc": 0.9}, "hw_available": ["h264_nvenc"],
                     "recommended": "h264_nvenc"})
assert video.resolve_codec("auto") == "h264_nvenc", "自动应挑实测最快的那个"
video.set_selfcheck(None)
assert video.resolve_codec("auto") == "libx264", "没自检过就保守用软编"
assert video.resolve_codec("hw") == "libx264", "没有可用硬编时 hw 也得回软编"
video.disable_hardware()
assert video.resolve_codec("h264_qsv") == "libx264", "运行期禁掉硬编后一律软编"
video._hw_disabled = False
print("编码器选择 OK（auto/hw/sw/具体名，未自检与禁用后都回软编）")


class _BadProc:
    returncode = 1
    stderr = "Error initializing MFX: 系统找不到指定的文件"


seg = OUT / "_seg_hw.mp4"


def _bad_run(cmd, **kw):
    captured.append(cmd)
    Path(cmd[-1]).unlink(missing_ok=True)     # 失败时不该留下半截文件
    return _BadProc()


video.subprocess.run = _bad_run
try:
    try:
        video.make_segment(OUT / "eq1.png", None, 3.0, seg, 1920, 1080, codec="h264_qsv")
        raise AssertionError("硬编失败应抛 EncoderFallback")
    except video.EncoderFallback as exc:
        assert "h264_qsv" in str(exc), exc
    assert "nv12" in captured[-1] and "-global_quality" in captured[-1], captured[-1]
    try:
        video.make_segment(OUT / "eq1.png", None, 3.0, seg, 1920, 1080)
        raise AssertionError("软编失败应抛 RuntimeError")
    except video.EncoderFallback:
        raise AssertionError("软编失败不该被当成硬编回退") from None
    except RuntimeError:
        pass
finally:
    video.subprocess.run = _orig_run
print("编码器参数 OK（硬编走 nv12/等效质量选项，失败分得清硬编与软编）")

# 自检：桩掉 subprocess，验证候选遍历、结果归类、cancel 与文案
probe_seen: list[str] = []


def _probe_run(cmd, **kw):
    codec = cmd[cmd.index("-c:v") + 1]
    if codec != "libx264" and codec != "h264_qsv":
        return _BadProc()                      # 本机没有 N 卡 / A 卡
    Path(cmd[-1]).write_bytes(b"fake mp4")
    return _FakeProc()


video.subprocess.run = _probe_run
try:
    res = video.self_check(seconds=1.0, progress=probe_seen.append)
    res2 = video.self_check(seconds=1.0, cancel=lambda: True)
finally:
    video.subprocess.run = _orig_run
assert set(res["timings"]) == {"libx264", "h264_qsv"}, res["timings"]
assert res["hw_available"] == ["h264_qsv"], res
assert res["recommended"] == min(res["timings"], key=lambda k: res["timings"][k])
assert "h264_nvenc" in res["errors"] and "h264_amf" in res["errors"], res["errors"]
assert len(probe_seen) == 4, probe_seen
assert "推荐" in video.selfcheck_summary(res), video.selfcheck_summary(res)
assert res2["cancelled"] and not res2["timings"], res2
assert "未自检" in video.selfcheck_summary(res2)
assert "没有可用" in video.selfcheck_summary({"timings": {"libx264": 1.0}})
print("编码自检 OK（候选遍历、不可用归类、推荐最快、可中断）")

# 硬编在真实片段上失败 → 整节改回软编重做（concat 用 -c copy，混编码器会拼出坏文件）
fb = OUT / "_enc_fallback"
shutil.rmtree(fb, ignore_errors=True)
(fb / "slides").mkdir(parents=True)
(fb / "work").mkdir(parents=True)
for p in (1, 2):
    (fb / "slides" / f"slide_{p:04d}.png").write_bytes(b"png")

calls: list[tuple[str, str]] = []
logs: list[str] = []
concats: list[list[str]] = []
_orig_make, _orig_concat = video.make_segment, video.concat_segments
video.concat_segments = lambda segs, out: concats.append([Path(s).name for s in segs])


def _mk(image, audio, dur, seg_path, w, h, fps=10, crf=22, preset="veryfast",
        codec="libx264", hw_ok=False):
    calls.append((Path(seg_path).name, codec))
    if video.is_hw(codec) and not hw_ok:
        raise video.EncoderFallback(f"{codec} 初始化失败")
    Path(seg_path).write_bytes(b"mp4")


try:
    video.make_segment = lambda *a, **k: _mk(*a, **k)
    pl = pipeline.Pipeline(settings=pipeline.Settings(codec="h264_qsv", workers=1),
                           log=logs.append)
    pl._encode_range(fb / "slides", ["", ""], 1, 2, 1920, 1080, fb / "cache", None,
                     fb / "work", fb / "out.mp4", 2, None)
    assert calls == [("seg_0001.mp4", "h264_qsv"),
                     ("seg_0001.mp4", "libx264"),
                     ("seg_0002.mp4", "libx264")], calls
    assert concats == [["seg_0001.mp4", "seg_0002.mp4"]], concats
    assert any("改回 CPU 软编" in m for m in logs), logs
    assert any(m.startswith("视频编码：") for m in logs), logs

    # 禁用是全进程生效的：后面几节不再重复踩同一个坑
    calls.clear()
    concats.clear()
    video.make_segment = lambda *a, **k: _mk(*a, **k, hw_ok=True)   # 就算硬编其实能用
    pipeline.Pipeline(settings=pipeline.Settings(codec="h264_qsv", workers=1),
                      log=logs.append)._encode_range(
        fb / "slides", ["", ""], 1, 2, 1920, 1080, fb / "cache", None,
        fb / "work", fb / "out.mp4", 2, None)
    assert [c for _, c in calls] == ["libx264", "libx264"], calls

    calls.clear()
    logs.clear()
    concats.clear()
    video._hw_disabled = False      # 上一轮失败已把硬编禁了，这里要测"硬编正常"的情况
    video.make_segment = lambda *a, **k: _mk(*a, **k, hw_ok=True)
    pl2 = pipeline.Pipeline(settings=pipeline.Settings(codec="h264_qsv", workers=2),
                            log=logs.append)
    pl2._encode_range(fb / "slides", ["", ""], 1, 2, 1920, 1080, fb / "cache", None,
                      fb / "work", fb / "out.mp4", 2, None)
    assert [c for _, c in calls] == ["h264_qsv", "h264_qsv"], calls
    assert not any("改回" in m for m in logs), logs
finally:
    video.make_segment, video.concat_segments = _orig_make, _orig_concat
    video._hw_disabled = False
    video._announced.clear()
    shutil.rmtree(fb, ignore_errors=True)
print("硬编失败回退 OK（整节改软编重做，硬编正常时不回退）")

# ---------------- 子进程不弹命令行黑框 ----------------
# 打包成无控制台 exe 后，每个 ffmpeg / powershell 子进程都会自己分配一个新控制台，
# 一节视频十几个片段就是一串黑框。所有调用点都必须带 CREATE_NO_WINDOW。
import subprocess as _sp                      # noqa: E402

from kouboppt import docgen, ppt_render, subproc   # noqa: E402

if sys.platform != "win32":
    assert subproc.NO_WINDOW == {}, subproc.NO_WINDOW
else:
    assert subproc.NO_WINDOW == {"creationflags": _sp.CREATE_NO_WINDOW}, subproc.NO_WINDOW

    run_kw: list[dict] = []

    class _P:
        returncode = 0
        stderr = "Duration: 00:00:03.20, start: 0.000000"

    def _kw_run(cmd, **kw):
        run_kw.append(kw)
        if str(cmd[-1]).endswith(".mp4"):
            Path(cmd[-1]).write_bytes(b"fake mp4")
        if "-PdfPath" in cmd:
            Path(cmd[cmd.index("-PdfPath") + 1]).write_bytes(b"%PDF")
        return _P()

    nw = OUT / "_nowindow"
    shutil.rmtree(nw, ignore_errors=True)
    (nw / "slides").mkdir(parents=True)
    (nw / "a.mp3").write_bytes(b"mp3")          # 假音频：mutagen 读不了 → 走 ffmpeg 探测

    _orig_popen = _sp.Popen

    class _FakePopen:
        def __init__(self, cmd, **kw):
            run_kw.append(kw)
            self.stdout = iter(["APP KWPP.Application\n", "OK 1\n", "OK 2\n", "DONE\n"])
            self.stderr = iter([])
            self._polls = 0

        def poll(self):
            self._polls += 1
            return None if self._polls < 3 else 0     # 让看门狗线程能退出

        def wait(self, timeout=None):
            return 0

        def kill(self):
            pass

    # 三个模块都是 `import subprocess`，改的是同一个模块属性，最后一并还原
    _sp.run, _sp.Popen = _kw_run, _FakePopen
    try:
        video.make_segment(OUT / "eq1.png", None, 2.0, nw / "seg.mp4", 1920, 1080)
        video.concat_segments([nw / "seg.mp4", nw / "seg.mp4"], nw / "all.mp4")
        video.audio_duration(nw / "a.mp3")
        video.self_check(seconds=0.2, deadline=5.0)
        docgen.docx_to_pdf(nw / "笔记.docx", nw / "笔记.pdf")
        ppt_render.export_slides(nw / "x.pptx", nw / "slides", 1920, 1080, 1, 2, workdir=nw)
    finally:
        _sp.run, _sp.Popen = _orig_run, _orig_popen
        shutil.rmtree(nw, ignore_errors=True)
    assert len(run_kw) >= 6, run_kw
    assert all(k.get("creationflags") == _sp.CREATE_NO_WINDOW for k in run_kw), run_kw
    print(f"子进程不弹黑框 OK（{len(run_kw)} 次 ffmpeg / powershell 调用都带 CREATE_NO_WINDOW）")

# ---------------- 课节目录页码范围登记（教材原文.md 缓存有效性校验） ----------------
from kouboppt import courseware as cw
_rng = OUT / "_range_test.json"
cw._write_range(_rng, 12, 30)
assert cw._read_range(_rng) == (12, 30), cw._read_range(_rng)
assert cw._read_range(OUT / "no_such.json") is None            # 缺失→None（旧版产物按信任处理）
_rng.write_text("{坏", encoding="utf-8")
assert cw._read_range(_rng) is None                             # 损坏→None
_rng.write_text(json.dumps({"start": 5, "end": 2}), encoding="utf-8")
assert cw._read_range(_rng) is None                             # 非法区间→None
_rng.unlink()
print("页码范围登记 OK（_read_range/_write_range 与缓存校验）")

# ---------------- XML 非法字符：写 Word/PPT 不能再被控制字符炸掉 ----------------
# 模型偶尔会在题干/公式里混入控制字符，python-docx 会抛
# ValueError: All strings must be XML compatible... 并中止整轮任务（实测出现过）。
from kouboppt import xmlsafe

assert xmlsafe.clean("a\x0bb") == "ab", repr(xmlsafe.clean("a\x0bb"))
assert xmlsafe.clean("a\x00b\x1fc") == "abc"
assert xmlsafe.clean("a\tb\nc\rd") == "a\tb\nc\rd", "\\t \\n \\r 必须保留"
assert xmlsafe.clean("正常文本") == "正常文本" and xmlsafe.clean("") == ""
assert xmlsafe.illegal("a\x0bb\x00c\x0b") == ["U+0000", "U+000B"], \
    xmlsafe.illegal("a\x0bb\x00c\x0b")
assert xmlsafe.illegal("clean") == []

BOOM, NULL = "\x0b", "\x00"          # 垂直制表符 / NUL：XML 1.0 都不允许

# 题库：题干/选项/答案/解析，且控制字符同时出现在公式段与普通文字段
bad_qs = [
    {"type": "single",
     "stem": f"已知 $a{BOOM}b$ 与 $c$ 的关系{NULL}如下",
     "options": ["A. $\\frac1n$", f"B. {BOOM}$\\sqrt2$"],
     "answer": f"A{BOOM}",
     "explanation": f"因为 $x\\le 0${NULL}，所以选 A"},
    {"type": "blank", "stem": f"填空：{BOOM}$E=mc^2${NULL}", "options": [],
     "answer": "略", "explanation": NULL},
]
q_path = docgen.questions_to_docx(bad_qs, OUT / "xmlsafe_quiz.docx", "测试教材", "第1课",
                                  with_answers=True, formula_dir=OUT / "xmlsafe_eq")
assert q_path.exists() and q_path.stat().st_size > 0, "含控制字符的题库应能正常落盘"
_qdoc = _Doc(q_path)
_qtxt = "\n".join(p.text for p in _qdoc.paragraphs)
assert "已知" in _qtxt and BOOM not in _qtxt and NULL not in _qtxt, repr(_qtxt[:200])

# 笔记 / 课后习题（markdown 路径：标题、列表、表格都要覆盖）
bad_md = (f"# 第1课{BOOM} 学习笔记\n\n"
          f"- 定义：$\\lim_{{x\\to 0}} f(x) = A${NULL}\n\n"
          f"| 量{BOOM} | 值 |\n|---|---|\n| $x$ | 1{NULL} |\n")
md_path = docgen.markdown_to_docx(bad_md, OUT / "xmlsafe_note.docx", "测试教材", "第1课",
                                  formula_dir=OUT / "xmlsafe_eq")
assert md_path.exists() and md_path.stat().st_size > 0, "含控制字符的笔记应能正常落盘"

# PPT：封面标题、小节标题、页标题、正文、备注（口播稿）全都要过
bad_lesson = slidegen.LessonSpec(
    title=f"第1章{BOOM}控制字符",
    subtitle=f"副标题{NULL}",
    sections=[slidegen.SectionSpec(
        title=f"小节{BOOM}",
        slides=[slidegen.SlideSpec(
            title=f"页标题{BOOM}",
            lines=[f"正文{NULL}里有 {BOOM}$x\\le 0$ 公式", "$$\\frac1n$$"],
            script=f"口播稿{BOOM}要能写进备注{NULL}。",
        )],
    )],
)
bad_pptx = slidegen.build(bad_lesson, OUT / "xmlsafe_lesson.pptx", theme="stem_blue")
assert bad_pptx.exists() and bad_pptx.stat().st_size > 0, "含控制字符的 PPT 应能正常生成"

# 万一仍然写盘失败：日志要给出一行摘要 + 命中的非法码位，并把源数据 repr 留档
def _boom():
    raise ValueError("模拟写盘失败")

_logs: list[str] = []
try:
    cw._doc_or_report("测试文档", OUT / "_xmlsafe_report", "含\x0b的源数据", _boom, _logs.append)
    assert False, "应当把异常抛出去"
except ValueError:
    pass
_lt = "\n".join(_logs)
assert "测试文档 生成失败" in _lt and "U+000B" in _lt, _lt
assert (OUT / "_xmlsafe_report" / "AI原始回复.log").exists(), "失败时应留档原始数据"
shutil.rmtree(OUT / "_xmlsafe_report", ignore_errors=True)
print("XML 非法字符清洗 OK（题库/笔记/PPT 不再被控制字符炸掉；失败有留档）")

# ---------------- 长度截断：原始回复也要留档（不能只在 JSON 坏掉时才留） ----------------
from kouboppt import llm as _llm


class _FakeResp:
    status_code = 200
    text = ""

    def __init__(self, reason: str, content: str):
        self._reason, self._content = reason, content

    def json(self):
        return {"choices": [{"finish_reason": self._reason,
                             "message": {"content": self._content}}]}


class _FakeSession:
    def __init__(self, reason: str, content: str):
        self._reason, self._content = reason, content

    def post(self, *a, **kw):
        return _FakeResp(self._reason, self._content)


_partial = '{"slides":[{"title":"半截内容'
_dump_dir = OUT / "_trunc_test"
shutil.rmtree(_dump_dir, ignore_errors=True)
_orig_session = _llm._session
try:
    _llm._session = lambda: _FakeSession("length", _partial)      # 模拟"撞到输出上限"
    _logs: list[str] = []
    _client = _llm.LLMClient(_llm.LLMConfig(base_url="http://127.0.0.1/v1"), log=_logs.append)
    assert _client.chat("s", "u", raw_dir=_dump_dir, tag="题库第1-8题") == _partial
    assert any("finish_reason=length" in s for s in _logs), _logs
    _dumped = (_dump_dir / "AI原始回复.log").read_text(encoding="utf-8")
    assert "题库第1-8题（长度截断）" in _dumped, _dumped        # 标签注明是长度截断
    assert _partial in _dumped, _dumped

    # 正常结束（stop）不该留档，免得日志目录被正常回复塞满
    shutil.rmtree(_dump_dir, ignore_errors=True)
    _llm._session = lambda: _FakeSession("stop", '{"slides":[]}')
    _logs2: list[str] = []
    _c2 = _llm.LLMClient(_llm.LLMConfig(base_url="http://127.0.0.1/v1"), log=_logs2.append)
    _c2.chat("s", "u", raw_dir=_dump_dir, tag="正常")
    assert not _logs2, _logs2
    assert not (_dump_dir / "AI原始回复.log").exists(), "正常回复不该留档"
finally:
    _llm._session = _orig_session
    shutil.rmtree(_dump_dir, ignore_errors=True)
print("长度截断留档 OK（length 留档并注明；stop 不留）")

# ---------------- 裸上标/幂：自动包成公式（否则 ^ 会原样显示成尖号） ----------------
# 模型在中文句子里写科学计数法时常忘掉 $（如 "5.00×10^22 个"），实测题库里就是这样。
assert formula.has_math("A. 5.00×10^22 个"), "裸科学计数法该被识别为公式"
assert formula.has_math("面积 3m^2"), "单位幂该被识别为公式"
assert formula.has_math("速度 v^2"), "变量幂该被识别为公式"
assert not formula.has_math("第一章^第2节"), "中文里的 ^ 不该被当公式"

_w = formula.normalize_delims
assert "$5.00\\times10^{22}$" in _w("A. 5.00×10^22 个"), _w("A. 5.00×10^22 个")
assert "\\times" in _w("B. 4.42x10^22 个"), _w("B. 4.42x10^22 个")
assert "$2.50\\times10^{-3}$" in _w("C. 2.50×10^{-3} mol"), _w("C. 2.50×10^{-3} mol")
assert "$1.00\\times10^{23}$" in _w("D. 1.00×10^23 个"), _w("D. 1.00×10^23 个")
assert "$10^{23}$" in _w("个数约为 10^23 个"), _w("个数约为 10^23 个")   # 无系数的裸 10^N
assert "$m^{2}$" in _w("面积 3m^2"), _w("面积 3m^2")            # 单字母按变量（斜体）
assert "$\\mathrm{cm}^{3}$" in _w("体积 3cm^3"), _w("体积 3cm^3")  # 多字母按单位（正体）
assert "$v^{2}$" in _w("速度 v^2"), _w("速度 v^2")
assert _w("$a^{2}$ 已经是公式") == "$a^{2}$ 已经是公式", _w("$a^{2}$ 已经是公式")
assert _w("第一章^第2节") == "第一章^第2节", _w("第一章^第2节")
assert _w("没有幂") == "没有幂"

# 已经写了 $ 但指数没加花括号：mathtext 里 x^22 只作用于后一个 token（会渲染成 x²2）
assert formula.sanitize_latex(r"$x^22$") == r"$x^{22}$", formula.sanitize_latex(r"$x^22$")
assert formula.sanitize_latex(r"$10^-3$") == r"$10^{-3}$", formula.sanitize_latex(r"$10^-3$")
assert formula.sanitize_latex(r"$x^2$") == r"$x^2$", "单位数指数不用动"
assert formula.sanitize_latex(r"$\mathrm{m}^{2}$") == r"$\mathrm{m}^{2}$", "已带括号不动"

# 端到端：题库选项里的 10^22 变成公式图片，不再留裸尖号
_pow_qs = [{"type": "single", "stem": "一杯水中水分子的个数约为？",
            "options": ["A. 5.00×10^22 个", "B. 4.42×10^22 个"],
            "answer": "A", "explanation": "由 $n=N/N_A$ 得，约 5.00×10^22。"}]
_pow_path = docgen.questions_to_docx(_pow_qs, OUT / "pow_quiz.docx", "测试教材", "第1课",
                                     with_answers=True, formula_dir=OUT / "pow_eq")
_pow_doc = _Doc(_pow_path)
_pow_txt = "\n".join(p.text for p in _pow_doc.paragraphs)
assert "^" not in _pow_txt, f"不该再有裸上标：{_pow_txt!r}"
assert len(_pow_doc.inline_shapes) >= 2, "幂应渲染成公式图片"
print("裸上标自动公式化 OK（10^22 / m^2 / v^2 自动包 $；已有的 $ 不动；x^22 补括号）")

# ---------------- 嵌套 \text/\mathrm：\frac{\mathrm{d}y}{\mathrm{d}x} 这类必须能渲染 ----------------
# 旧实现按 \text/\mathrm 切分公式，命令嵌在 {} 组里时会切出 "\frac{" 这种未闭合残段，
# 整行渲染失败 → 降级把 LaTeX 源码写进文档（题库与 PPT 里都实测出现过）。
_nested = [
    r"$a_{\text{Si}}$", r"$a_{\text{Ge}}$", r"$x_{\text{max}}$", r"$v^{\text{out}}$",
    r"$\frac{\mathrm{d}E_g}{\mathrm{d}P}$", r"$\frac{\mathrm{d}y}{\mathrm{d}x}$",
    r"$$\frac{\mathrm{d}E_g}{\mathrm{d}P}$$", r"$\frac{\mathrm{d}N}{\mathrm{d}V}$",
    r"禁带宽度随压强的变化：$$\frac{\mathrm{d}E_g}{\mathrm{d}P}=-(\Xi_c-\Xi_v)$$",
]
for _i, _tex in enumerate(_nested):
    assert formula.render(_tex, OUT / f"nest{_i}.png",
                          display=_tex.strip().startswith("$$")), \
        f"嵌套 \\text/\\mathrm 仍渲染失败：{_tex}"

# 切分不能再产出残缺段（$\frac{$ 这种）
assert formula._math_body(r"a_{\text{Si}}") == r"$a_{\mathrm{Si}}$", \
    formula._math_body(r"a_{\text{Si}}")
assert formula._math_body(r"\frac{\mathrm{d}E_g}{\mathrm{d}P}") == \
    r"$\frac{\mathrm{d}E_g}{\mathrm{d}P}$", formula._math_body(r"\frac{\mathrm{d}E_g}{\mathrm{d}P}")
# 顶层 \text 仍要提到 math 之外（中文用中文字体渲染）
assert formula._math_body(r"\text{速度}") == "速度"
assert formula._math_body(r"\mathrm{d}x") == r"$\mathrm{d}$$x$"
# 中文留在 math 里时不再返回 None：切到中文字体集渲染（`E_{\text{总}}` 因此能出图）
assert formula._math_body(r"v_{\text{初}}") == r"$v_{\mathrm{初}}$", \
    formula._math_body(r"v_{\text{初}}")
assert formula.render(r"$v_{\text{初}}$", OUT / "cjk_ok.png") is not None, \
    "含中文下标的公式现在应能渲染"

# to_readable：渲染失败时的可读降级，而不是把 LaTeX 源码甩给用户
assert formula.to_readable(r"$a_{\text{Si}} = 0.543102\ \text{nm}$") == "a_Si = 0.543102 nm", \
    formula.to_readable(r"$a_{\text{Si}} = 0.543102\ \text{nm}$")
assert formula.to_readable(r"\frac{\mathrm{d}E_g}{\mathrm{d}P}") == "(dE_g)/(dP)"
assert formula.to_readable(r"$5.00\times10^{22}$") == "5.00×10²²"
assert formula.to_readable(r"$v_{\text{初}}$") == "v_初"
assert formula.to_readable(r"$x^2+y^2=r^2$") == "x²+y²=r²"
assert "$" not in formula.to_readable(r"前言 $$x=1$$ 后记")

# 端到端：含中文的公式现在出图；真渲染不了的才落成可读文本，而且必须被收集留档
docgen.reset_failed()
_read_qs = [{"type": "short", "stem": r"求 $v_{\text{初}}$ 的值", "options": [],
             "answer": "略",
             "explanation": r"画不出的：$\qqqbadcmd{x}$；能画的：$a_{\text{Si}}$"}]
_read_path = docgen.questions_to_docx(_read_qs, OUT / "readable_quiz.docx", "测试教材",
                                      "第1课", with_answers=True,
                                      formula_dir=OUT / "readable_eq")
_read_fails = docgen.drain_failed()
_read_txt = "\n".join(p.text for p in _Doc(_read_path).paragraphs)
assert "\\text" not in _read_txt and "\\mathrm" not in _read_txt, \
    f"不该出现 LaTeX 源码：{_read_txt!r}"
assert "v_初" not in _read_txt, f"含中文下标的公式现在应出图，不该降级成纯文本：{_read_txt!r}"
assert "qqqbadcmdx" in _read_txt, _read_txt
# Word 侧过去是静默降级（用户只看到纯文本，日志里毫无提示），现在必须收集得到
assert len(_read_fails) == 1 and "qqqbadcmd" in _read_fails[0], _read_fails
assert docgen.drain_failed() == [], "drain 之后应清空"

# 失败清单单独留档（日志里只够显示一条）
_fail_dir = OUT / "_fail_formula"
shutil.rmtree(_fail_dir, ignore_errors=True)
cw._dump_failed_formulas(_fail_dir, [r"\frac{\mathrm{d}E_g}{\mathrm{d}P}", "$v_{\text{初}}$"],
                         lambda s: None)
_fail_log = (_fail_dir / "公式渲染失败.log").read_text(encoding="utf-8")
assert "共 2 处" in _fail_log and "mathrm" in _fail_log, _fail_log
shutil.rmtree(_fail_dir, ignore_errors=True)
print("嵌套 \\text/\\mathrm + 可读降级 OK（导数式可渲染；失败降级为可读文本并留档）")

print("ALL OK")
