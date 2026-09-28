"""离线端到端测试：桩 LLM + 自造 PDF，验证 教材→课程 两步编排（不起视频渲染）。

覆盖：
1. 全自动模式一口气跑完（按每节页数切分）；
2. 分步模式：课件+视频做完后选「就到这里」→ 不出笔记，且 2 节的书只问一次；
3. 断点续跑：接着已有产物出笔记，不重复生成 PPT；
4. 只勾「学习笔记」→ 不生成 PPT，但仍识别原文；
5. 按教材章节：AI 识别目录 → 章节结构.json 缓存 → 二次运行不再调 AI；
6. 章节内细分不跨章、每页只出现一次；
7. 并发正确性：识别按页并发且结果与串行一致、小节并发、视频按页并发，
   workers=1 时完全退回串行；
8. 拆小块防截断：内容页 ≤3 页/次、题库 ≤8 题/次，拼起来完整有序；
9. 截断抢救：坏 JSON 里已写完整的内容救下来继续用，原始回复留档；
10. 课件/视频流水线：上一节视频合成时下一节课件已在生成，视频仍按书序出；
11. 视频收尾时机：全自动下笔记/题库与后台视频并行、分步确认仍在闸门前收完，
    第二步出错时后台视频立刻收到停止信号；
12. 产物分文件夹：课件 → PPT/、视频 → 视频/，文件名 = 课时标题 + 类型；
    旧结构（都堆在课节目录里）自动搬迁进新文件夹，搬完即算缓存。
13. 课后习题/思考题：从识别稿里切出来单独出 Word（课后习题/课时标题课后习题.docx），
    正文（不含习题）才用于课件/笔记/题库，重跑算缓存。
"""
import re
import shutil
import sys
import threading
import time
from pathlib import Path

import pymupdf

# 控制台是 GBK 时打印 ✔/⚠ 会 UnicodeEncodeError，导致测试整个失败
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(errors="replace")


class Overlap:
    """记录"同一时刻在跑的调用数"峰值，用来证明真的并发了。"""

    def __init__(self):
        self.lock = threading.Lock()
        self.running = 0
        self.peak = 0

    def __enter__(self):
        with self.lock:
            self.running += 1
            self.peak = max(self.peak, self.running)

    def __exit__(self, *exc):
        with self.lock:
            self.running -= 1


class StubLLM:
    """按提示词特征返回固定内容，代替真实大模型；记录调用类型与收到的提示词。"""

    def __init__(self):
        self.calls: list[str] = []
        self.sent: list[str] = []          # 每条请求的 user 提示词，供"是否带上教材原文"检查
        self.log = lambda s: None          # 编排层会通过 llm.log 报"抢救/截断"这类警告

    def chat(self, system: str, user: str, images=None, max_tokens: int = 0,
             **kw) -> str:
        self.sent.append(user)
        if "转写为 Markdown" in user:
            self.calls.append("ocr")
            return "# 标题\n正文 $a^2+b^2=c^2$。\n\n$$\\int_0^1 x\\,dx = \\frac{1}{2}$$"
        if "学习笔记" in user:
            self.calls.append("notes")
            return ("# 知识框架\n## 核心概念\n- 勾股定理：$a^2+b^2=c^2$\n"
                    "## 重要公式\n$$\\oint_C \\vec{F}\\cdot d\\vec{r} = 0$$\n"
                    "- 注意：$\\varepsilon > 0$")
        raise AssertionError("桩 LLM 没预料到的 chat 调用：" + user[:80])

    def chat_json(self, system: str, user: str, images=None, max_tokens: int = 0,
                  **kw):
        self.sent.append(user)
        if "可选风格" in user:                   # 先判这句：选风格的提示词里也含"出…题"字样
            self.calls.append("style")
            return {"styles": [{"key": "science_green", "reason": "理科教材"},
                               {"key": "medical_teal", "reason": "备选一套"}]}
        if "推断每一章" in user:
            self.calls.append("chapters")
            return {"chapters": [{"title": "第1章 测试章节", "start": 1, "end": 2},
                                 {"title": "第2章 测试章节", "start": 3, "end": 4}]}
        if "大纲" in user:
            self.calls.append("outline")
            return {"sections": [
                {"title": "第一节 概念", "focus": "定义", "slides": 2},
                {"title": "第二节 例题", "focus": "计算", "slides": 2},
            ]}
        if "内容页" in user:
            self.calls.append("slides")
            return {"slides": [
                {"title": "勾股定理", "lines": [
                    "直角三角形两直角边平方和等于斜边平方",
                    "$$a^2 + b^2 = c^2$$",
                    "当 $c$ 为斜边时成立"],
                 "script": "同学们，今天我们学习勾股定理，a 的平方加 b 的平方等于 c 的平方。"},
                {"title": "例题", "lines": [
                    "已知 a=3，b=4，求 c",
                    "$$c = \\sqrt{a^2+b^2} = 5$$"],
                 "script": "我们来看一道例题，三的平方加四的平方等于二十五，开方得五。"},
            ]}
        if "出" in user and "题" in user:
            self.calls.append("quiz")
            return {"questions": [
                {"type": "single", "stem": "勾股定理适用于？$a^2+b^2=c^2$",
                 "options": ["A. 任意三角形", "B. 直角三角形", "C. 等边三角形"],
                 "answer": "B", "explanation": "仅直角三角形两直角边满足。"},
                {"type": "blank", "stem": "$\\sqrt{3^2+4^2}=$ ____", "options": [],
                 "answer": "$5$", "explanation": "3-4-5 勾股数。"},
            ]}
        raise AssertionError("桩 LLM 没预料到的 chat_json 调用：" + user[:80])


class SlowLLM(StubLLM):
    """每次调用都睡一小会儿并记录并发峰值的桩，用于验证并发路径。"""

    def __init__(self, delay: float = 0.15):
        super().__init__()
        self.delay = delay
        self.ov = Overlap()
        self.page_order: list[int] = []

    def chat(self, system: str, user: str, images=None, max_tokens: int = 0,
             **kw) -> str:
        with self.ov:
            time.sleep(self.delay)
            m = re.search(r"第 (\d+) 页", user)
            page = int(m.group(1)) if m else 0
            self.page_order.append(page)
            self.calls.append("ocr")
            return f"# 第 {page} 页正文\n这是第 {page} 页的 $x_{{{page}}}$。"

    def chat_json(self, system: str, user: str, images=None, max_tokens: int = 0,
                  **kw):
        with self.ov:
            time.sleep(self.delay)
            if "大纲" in user:
                self.calls.append("outline")
                return {"sections": [{"title": f"第{i}节", "focus": "x", "slides": 1}
                                     for i in range(1, 5)]}
            if "内容页" in user:
                self.calls.append("slides")
                title = re.search(r"小节《(.+?)》", user).group(1)
                return {"slides": [{"title": title, "lines": ["要点一", "要点二"],
                                    "script": f"{title} 的口播稿。"}]}
            raise AssertionError("桩 LLM 没预料到的 chat_json 调用：" + user[:80])


class BadFormulaLLM(StubLLM):
    """内容里塞一条 mathtext 渲不出的公式（同位素写法），并应答兜底的两类请求。"""

    BAD = r"碳的同位素写法：$\prescript{14}{6}{C}$"

    def chat_json(self, system: str, user: str, images=None, max_tokens: int = 0, **kw):
        if "只改写法" in user:                   # 兜底层①：公式改写
            self.calls.append("fix_formula")
            return {"latex": r"^{14}_{6}\mathrm{C}", "note": "换成 mathtext 写法"}
        out = super().chat_json(system, user, images=images, max_tokens=max_tokens, **kw)
        if isinstance(out, dict) and out.get("slides"):    # 每页内容都换成这条坏公式
            for slide in out["slides"]:
                slide["lines"] = [self.BAD]
        return out


def make_pdf(path: Path, chapter_pages=(1, 1)):
    """每章按 chapter_pages 里的数字造页（默认各 1 页 → 2 页 PDF）。"""
    doc = pymupdf.open()
    for ch, n_pages in enumerate(chapter_pages, 1):
        for _ in range(n_pages):
            text = "\n".join([
                f"第{ch}章 测试章节",
                "直角三角形的两条直角边的平方和等于斜边的平方。",
                "这就是著名的勾股定理，在古代中国被称为商高定理。",
                "典型应用是已知两边求第三边，例如 3、4、5。",
                "本章还会介绍正弦定理与余弦定理的推导过程和应用。",
            ] * 6)
            page = doc.new_page()
            page.insert_textbox(pymupdf.Rect(72, 72, 520, 760), text, fontsize=12,
                                fontname="china-s")
    doc.save(path)
    doc.close()


def fresh_dir(d: Path) -> Path:
    if d.exists():
        shutil.rmtree(d)
    d.mkdir(parents=True, exist_ok=True)
    return d


def make_scanned_pdf(path: Path, pages: int) -> None:
    """文字层里带 $ → 走多模态识别路径（模拟扫描页）。"""
    doc = pymupdf.open()
    for i in range(1, pages + 1):
        page = doc.new_page()
        page.insert_textbox(pymupdf.Rect(72, 72, 520, 760),
                            f"第 {i} 页 扫描内容 $x_{{{i}}}$", fontsize=12)
    doc.save(path)
    doc.close()


def make_opts(out_dir: Path, chapter_pages=(1, 1), **kw):
    from kouboppt import courseware
    from kouboppt.pipeline import Settings

    pdf = out_dir / "测试教材.pdf"
    make_pdf(pdf, chapter_pages)
    base = dict(pdf_path=pdf, out_dir=out_dir, split_mode="pages",
                start_page=1, pages_per_lesson=1,
                target_minutes=35, gen_ppt=True, gen_video=False,
                gen_notes=True, gen_quiz=True, quiz_count=4, video=Settings())
    base.update(kw)
    return courseware.Options(**base)


def log_quiet(s: str):
    if s.strip().startswith(("✔", "════", "✘", "⚠", "⏸", "全部", "已停")):
        print("  " + s)


def idx_of(name: str) -> int:
    """本测试的课时标题形如「第1课 p1-p1」：从标题取节号，不再依赖文件放在哪个目录。"""
    m = re.search(r"第(\d+)", str(name))
    assert m, f"标题里没有节号：{name}"
    return int(m.group(1))


def test_auto():
    from kouboppt import courseware

    out = fresh_dir(Path("test_out/course_auto"))
    opts = make_opts(out)
    llm = StubLLM()
    outs = courseware.run(llm, opts, log=log_quiet, progress=lambda f, m: None)

    names = {p.name for p in outs}
    ppt = next(p for p in outs if p.name.endswith("PPT.pptx"))
    assert ppt.parent.name == courseware.PPT_DIR, f"课件该在 {courseware.PPT_DIR}/ 里：{ppt}"
    assert any("学习笔记.docx" in n for n in names), "缺少笔记"
    assert any("题库.docx" in n for n in names), "缺少题库"
    assert len(list(out.rglob("教材原文.md"))) == 2, "两章都该有识别文本"
    assert llm.calls.count("outline") == 2, llm.calls
    assert llm.calls.count("notes") == 4, f"每章笔记分上下篇两次请求：{llm.calls}"
    assert "gate" not in llm.calls
    print("  → 全自动模式 OK（产出", len(outs), "个文件）")


def test_auto_style():
    """theme="auto"：按书名让 AI 选风格，落盘 PPT风格.json（全书只问一次），课件照常出。"""
    from kouboppt import courseware, slidegen

    out = fresh_dir(Path("test_out/course_style"))
    opts = make_opts(out, theme="auto", gen_notes=False, gen_quiz=False)
    llm = StubLLM()
    outs = courseware.run(llm, opts, log=log_quiet, progress=lambda f, m: None)

    assert llm.calls.count("style") == 1, f"选风格该整本书只问一次：{llm.calls}"
    files = list(out.rglob(slidegen.STYLE_FILE))
    assert files, "风格决定该落盘到书目录"
    assert slidegen.load_style_choice(files[0].parent)["key"] == "science_green", files
    assert any(p.name.endswith("PPT.pptx") for p in outs), "自动风格下也该出课件"
    print("  → 自动选风格 OK（只问一次、落盘缓存、课件照常出）")


def test_formula_fallback():
    """公式兜底两层：AI 改写重渲（层①）→ 扫图裁图（层②）→ 才降级；成功失败都落盘缓存。"""
    import pymupdf

    from kouboppt import formula, formula_fallback, textbook

    root = fresh_dir(Path("test_out/formula_fallback"))
    cache, pages = root / "cache", root / "pages"
    bad = r"$\prescript{14}{6}{C}$"
    assert formula.render(bad, root / "probe.png", fontsize=20, color="#000000",
                          display=False) is None, "样本必须本地渲不出"

    class RescueLLM:
        """只应答兜底请求：公式改写（层①）/ 公式定位（层②）。"""

        def __init__(self, rewritten="", box=None, index=1):
            self.rewritten, self.box, self.index, self.tags = rewritten, box, index, []

        def chat_json(self, system, user, images=None, max_tokens=0, **kw):
            self.tags.append(kw.get("tag"))
            if kw.get("tag") == "公式改写":
                return {"latex": self.rewritten}
            if kw.get("tag") == "公式定位":
                return {"index": self.index, "box": self.box}
            raise AssertionError("桩没预料到的请求：" + str(kw.get("tag")))

    # ---- 层①：AI 改写后本地能渲 → 直接用这张图；第二次命中缓存，不再问 AI
    fix = RescueLLM(rewritten=r"^{14}_{6}\mathrm{C}")
    fb = formula_fallback.make_fallback(fix, None, (1, 1), cache / "l1", pages,
                                        log=lambda s: None)
    img = fb(bad, "#000000", False)
    assert img is not None and img.exists() and img.stat().st_size > 0, "层①该给出一张图"
    assert img.name.startswith("fb_"), img.name
    assert fix.tags == ["公式改写"], fix.tags
    assert fb(bad, "#000000", False) == img and len(fix.tags) == 1, "第二次该命中缓存"

    # ---- 层②：改写救不回来 → 在教材原页上定位这条公式并裁图
    doc = pymupdf.open()
    page = doc.new_page()
    page.draw_rect(pymupdf.Rect(178, 315, 417, 345), color=None, fill=(0, 0, 0))
    doc.save(root / "书页.pdf")
    textbook.render_page(doc, 1, pages, dpi=100)          # 层②要用的页面图
    found = RescueLLM(rewritten="", box=[29, 36, 71, 42])  # 百分比框，正好圈住那个黑块
    fb2 = formula_fallback.make_fallback(found, doc, (1, 1), cache / "l2", pages,
                                         log=lambda s: None)
    crop = fb2(bad, "#000000", True)
    assert crop is not None and crop.exists() and crop.stat().st_size > 0, "层②该裁到一张图"
    assert crop.name.startswith("crop_"), crop.name
    assert found.tags == ["公式改写", "公式定位"], found.tags
    from PIL import Image
    with Image.open(crop) as im:
        assert im.height < 200, f"应只裁到那一小块而不是整页：{im.size}"
        assert im.convert("L").getextrema()[0] < 100, "裁出来的图里该有内容"

    # ---- 两级都失败：返回 None，并记住"救不了"，不再反复花钱
    dead = RescueLLM(rewritten="", box=[0, 0, 1, 1])       # 改写原样返回 + 框不合法
    fb3 = formula_fallback.make_fallback(dead, doc, (1, 1), cache / "l3", pages,
                                         log=lambda s: None)
    assert fb3(bad, "#000000", False) is None
    assert fb3(bad, "#000000", False) is None
    assert dead.tags.count("公式改写") == 1 and dead.tags.count("公式定位") == 1, \
        f"失败也该缓存，第二次不该再问：{dead.tags}"
    print("  → 公式兜底两层 OK（改写重渲 / 原页裁图 / 救不了也记住别再花钱）")


def test_formula_fallback_e2e():
    """端到端：课件里渲不出的公式靠兜底留在 PPT 里（不降级成文字），重跑吃缓存。"""
    from kouboppt import courseware

    out = fresh_dir(Path("test_out/course_formula"))
    opts = make_opts(out, chapter_pages=(1,), gen_notes=False, gen_quiz=False)
    llm = BadFormulaLLM()
    lines: list[str] = []
    outs = courseware.run(llm, opts, log=lines.append, progress=lambda f, m: None)

    assert llm.calls.count("fix_formula") == 1, f"每条失败公式问一次改写：{llm.calls}"
    assert any("兜底图补上" in s for s in lines), "\n".join(lines[-12:])
    assert not list(out.rglob("公式渲染失败.log")), "救回来了就不该留下失败清单"
    assert any(p.name.endswith("PPT.pptx") for p in outs), "课件照常出"

    # 删掉课件重做一遍：兜底结果有缓存，不该再问 AI
    for p in list(out.rglob("*PPT.pptx")) + list(out.rglob("口播稿.txt")):
        p.unlink()
    llm2 = BadFormulaLLM()
    courseware.run(llm2, opts, log=lambda s: None, progress=lambda f, m: None)
    assert llm2.calls.count("fix_formula") == 0, f"兜底缓存该生效：{llm2.calls}"
    print("  → 公式兜底端到端 OK（坏公式进 PPT、重跑命中缓存、无失败清单）")


def test_prompts_carry_source():
    """内容生成的每一步都必须把教材原文喂给 AI——大纲/题库曾漏传原文，AI 自编数学内容。"""
    from kouboppt import courseware

    out = fresh_dir(Path("test_out/course_prompt"))
    opts = make_opts(out, chapter_pages=(1,))
    llm = StubLLM()
    courseware.run(llm, opts, log=log_quiet, progress=lambda f, m: None)

    for kw, name in (("大纲", "课件大纲"), ("内容页", "小节内容"), ("学习笔记", "学习笔记"),
                     ("道题", "题库")):
        hits = [u for u in llm.sent if kw in u]
        assert hits, f"没发出「{name}」请求"
        assert all("勾股定理" in u for u in hits), f"「{name}」请求里没带教材原文"
    print(f"  → 大纲/内容页/笔记/题库共 {len(llm.sent)} 条请求都带上了教材原文 OK")


def test_step_and_resume():
    from kouboppt import courseware

    out = fresh_dir(Path("test_out/course_step"))
    opts = make_opts(out, gen_quiz=False)          # 计划 = [第一步, 第三步]
    llm = StubLLM()
    asked: list[str] = []

    def gate(next_name, done, out_dir):
        asked.append(next_name)
        return False                               # 第一步完成后选「就到这里」

    outs = courseware.run(llm, opts, log=log_quiet, progress=lambda f, m: None,
                          gate=gate)
    assert len(asked) == 1, f"2 章的书应该只在第一步结束后问一次，实际 {asked}"
    assert "第二步" in asked[0], asked
    assert list(out.rglob("*PPT.pptx")), "第一步该有课件"
    assert not list(out.rglob("*学习笔记.docx")), "选「就到这里」后不该出笔记"
    assert llm.calls.count("notes") == 0, llm.calls
    assert any(p.name == "口播稿.txt" for p in outs), "缺少口播稿"
    print("  → 分步停下 OK（问询：", asked[0], "）")

    # 断点续跑：只勾「学习笔记」，接着已有产物继续
    before = {p: p.stat().st_mtime for p in out.rglob("*.pptx")}
    llm2 = StubLLM()
    opts2 = make_opts(out, gen_ppt=False, gen_quiz=False)
    outs2 = courseware.run(llm2, opts2, log=log_quiet, progress=lambda f, m: None)
    assert list(out.rglob("*学习笔记.docx")), "续跑该出笔记"
    assert llm2.calls.count("notes") == 4, llm2.calls
    assert llm2.calls.count("outline") == 0, f"课件该被复用而不是重做：{llm2.calls}"
    after = {p: p.stat().st_mtime for p in out.rglob("*.pptx")}
    assert before == after, "课件被重写了，应该原样复用"
    names2 = {p.name for p in outs2}
    assert any("学习笔记.docx" in n for n in names2), names2
    print("  → 断点续跑 OK（复用课件，未重复生成）")


def test_notes_only():
    from kouboppt import courseware

    out = fresh_dir(Path("test_out/course_notes_only"))
    opts = make_opts(out, gen_ppt=False, gen_quiz=False)
    llm = StubLLM()
    outs = courseware.run(llm, opts, log=log_quiet, progress=lambda f, m: None)
    assert not list(out.rglob("*.pptx")), "没勾 PPT 就不该生成课件"
    assert len(list(out.rglob("*学习笔记.docx"))) == 2, "两章都该有笔记"
    assert llm.calls.count("outline") == 0, llm.calls
    assert any("学习笔记.docx" in p.name for p in outs)
    print("  → 只出笔记 OK（未生成 PPT）")


def test_chapter_mode():
    from kouboppt import courseware
    import json

    out = fresh_dir(Path("test_out/course_chapters"))
    opts = make_opts(out, chapter_pages=(2, 2), split_mode="chapters",
                     pages_per_lesson=15, gen_quiz=False)
    llm = StubLLM()
    courseware.run(llm, opts, log=log_quiet, progress=lambda f, m: None)
    assert llm.calls.count("chapters") == 1, llm.calls
    assert len(list(out.rglob("教材原文.md"))) == 2, "每章 ≤15 页 → 每章一节"
    f = next(out.glob("*/章节结构.json"))
    data = json.loads(f.read_text(encoding="utf-8"))
    assert data["source"] == "ai" and len(data["chapters"]) == 2, data
    dirs = sorted(p.name for p in f.parent.glob("[0-9][0-9]_*") if p.is_dir())
    assert dirs == ["01_第1章_测试章节", "02_第2章_测试章节"], dirs

    llm2 = StubLLM()                        # 二跑：读缓存，不再调 AI
    courseware.run(llm2, opts, log=log_quiet, progress=lambda f, m: None)
    assert llm2.calls.count("chapters") == 0, llm2.calls
    assert llm2.calls.count("outline") == 0, llm2.calls
    print("  → 按教材章节 OK（AI 识别一次并缓存，二次运行零调用）")


def test_no_cross_chapter():
    from kouboppt import courseware

    out = fresh_dir(Path("test_out/course_chapters_split"))
    opts = make_opts(out, chapter_pages=(2, 2), split_mode="chapters",
                     pages_per_lesson=1)
    llm = StubLLM()
    courseware.run(llm, opts, log=log_quiet, progress=lambda f, m: None)
    mds = sorted(out.rglob("教材原文.md"))
    assert len(mds) == 4, [str(p) for p in mds]
    for i, f in enumerate(mds, 1):
        assert f"原书第 {i} 页" in f.read_text(encoding="utf-8"), f
    dirs = [m.parent.name for m in mds]
    assert dirs == ["01_第1章_测试章节_第1节", "02_第1章_测试章节_第2节",
                    "03_第2章_测试章节_第1节", "04_第2章_测试章节_第2节"], dirs
    print("  → 章节内细分 OK（每节 1 页，不跨章、不重不漏）")


def test_ocr_concurrency():
    from kouboppt import textbook

    out = fresh_dir(Path("test_out/conc_ocr"))
    pdf = out / "扫描教材.pdf"
    make_scanned_pdf(pdf, 8)
    doc = textbook.open_pdf(pdf)
    try:
        llm1 = SlowLLM()
        t0 = time.monotonic()
        md1 = textbook.ocr_chapter(llm1, doc, 1, 8, out / "img1",
                                   log=lambda s: None, workers=1)
        serial = time.monotonic() - t0
        assert llm1.ov.peak == 1, f"workers=1 应该完全串行，峰值 {llm1.ov.peak}"
        assert sorted(llm1.page_order) == list(range(1, 9)), llm1.page_order

        llm4 = SlowLLM()
        t0 = time.monotonic()
        md4 = textbook.ocr_chapter(llm4, doc, 1, 8, out / "img2",
                                   log=lambda s: None, workers=4)
        parallel = time.monotonic() - t0
        assert llm4.ov.peak >= 2, f"4 路并发没生效，峰值 {llm4.ov.peak}"
        assert md4 == md1, "并发识别结果应与串行逐字一致"
        pos = [md4.index(f"<!-- 原书第 {p} 页 -->") for p in range(1, 9)]
        assert pos == sorted(pos), f"页序被打乱：{pos}"
        assert parallel < serial * 0.8, f"并发没提速：串行 {serial:.2f}s / 并发 {parallel:.2f}s"
        assert not list((out / "img1").glob("*")) and not list((out / "img2").glob("*")), \
            "识别用的临时图片应被清理"
        print(f"  → 识别并发 OK（8 页：串行 {serial:.2f}s → 4 路 {parallel:.2f}s，"
              f"峰值 {llm4.ov.peak}；结果与串行一致、页序正确）")
    finally:
        doc.close()


def test_lesson_concurrency():
    from kouboppt import courseware

    llm1 = SlowLLM()
    one = courseware.gen_lesson(llm1, "教材原文", "第1课", "书", 35, "stem_blue", workers=1)
    assert llm1.ov.peak == 1, f"workers=1 应该完全串行，峰值 {llm1.ov.peak}"

    llm4 = SlowLLM()
    four = courseware.gen_lesson(llm4, "教材原文", "第1课", "书", 35, "stem_blue", workers=4)
    assert llm4.ov.peak >= 2, f"4 路并发没生效，峰值 {llm4.ov.peak}"
    titles = [sec.title for sec in four.sections]
    assert titles == ["第1节", "第2节", "第3节", "第4节"], titles
    assert [s.title for sec in four.sections for s in sec.slides] == \
           [s.title for sec in one.sections for s in sec.slides], "并发生成的页序应与串行一致"
    print(f"  → 小节并发 OK（4 小节峰值 {llm4.ov.peak}，小节顺序与内容不变）")


def test_notes_quiz_concurrency():
    from kouboppt import courseware

    class NotesQuizLLM(SlowLLM):
        def chat(self, system, user, images=None, max_tokens=0, **kw):
            with self.ov:
                time.sleep(self.delay)
                assert "学习笔记" in user, user[:60]
                return "# 知识框架\n- 要点 $a^2$\n$$\\frac{1}{2}$$"

        def chat_json(self, system, user, images=None, max_tokens=0, **kw):
            with self.ov:
                time.sleep(self.delay)
                return {"questions": [
                    {"type": "blank", "stem": "$1+1=$ ____", "options": [],
                     "answer": "$2$", "explanation": "略"}]}

    out = fresh_dir(Path("test_out/conc_notes"))
    opts = make_opts(out, chapter_pages=(1,), gen_ppt=False,
                     gen_notes=True, gen_quiz=True, quiz_count=4)
    llm = NotesQuizLLM()
    opts = courseware.Options(**{**vars(opts), "workers": 4})
    courseware.run(llm, opts, log=log_quiet, progress=lambda f, m: None)
    assert llm.ov.peak >= 2, f"笔记与题库没并发，峰值 {llm.ov.peak}"
    assert list(out.rglob("*学习笔记.docx")) and list(out.rglob("*题库.docx"))
    print(f"  → 笔记/题库并发 OK（同一节课两路请求峰值 {llm.ov.peak}，docx 正常产出）")


def test_video_page_concurrency():
    from kouboppt import pipeline, video

    out = fresh_dir(Path("test_out/conc_video"))
    slides = out / "slides"
    slides.mkdir(parents=True)
    n = 8
    for i in range(1, n + 1):
        (slides / f"slide_{i:04d}.png").write_bytes(b"png")

    class StubProvider:
        name = "stub"

        def synthesize(self, text, voice, rate, out_path):
            Path(out_path).write_bytes(b"mp3")

    concat_order: list[str] = []
    ov = Overlap()
    (out / "cache").mkdir(exist_ok=True)
    orig = (video.make_segment, video.audio_duration, video.concat_segments)

    def fake_make_segment(image, audio, duration, out_path, width, height, **kw):
        with ov:
            time.sleep(0.15)
            Path(out_path).write_bytes(b"seg")

    video.make_segment = fake_make_segment
    video.audio_duration = lambda p: 5.0
    video.concat_segments = lambda segs, out_path: concat_order.extend(
        [Path(s).name for s in segs])
    try:
        for workers, expect_peak in ((1, 1), (4, None)):
            concat_order.clear()
            ov.peak = 0
            pl = pipeline.Pipeline(pipeline.Settings(workers=workers), log=lambda s: None)
            workdir = out / f"w{workers}"
            workdir.mkdir(exist_ok=True)
            texts = [f"第 {i} 页文字。" for i in range(1, n + 1)]
            pl._encode_range(slides, texts, 1, n, 1920, 1080,
                             out / "cache", StubProvider(), workdir,
                             out / f"v{workers}.mp4", n, None)
            if expect_peak == 1:
                assert ov.peak == 1, f"workers=1 应该完全串行，峰值 {ov.peak}"
            else:
                assert ov.peak >= 2, f"4 路并发没生效，峰值 {ov.peak}"
            assert concat_order == [f"seg_{i:04d}.mp4" for i in range(1, n + 1)], concat_order
    finally:
        video.make_segment, video.audio_duration, video.concat_segments = orig
    print(f"  → 视频页级并发 OK（8 页峰值 {ov.peak}，拼接顺序仍为页序）")


def test_chunked_requests():
    """内容页按 ≤3 页/次、题库按 ≤8 题/次 拆小块请求（防接口截断），拼起来完整且顺序不乱。"""
    from kouboppt import courseware

    class ChunkLLM(StubLLM):
        def chat_json(self, system, user, images=None, max_tokens=0, **kw):
            if "大纲" in user:
                self.sent.append(user)
                self.calls.append("outline")
                return {"sections": [{"title": "第一节 概念", "focus": "定义", "slides": 7}]}
            if "内容页" in user:
                self.sent.append(user)
                self.calls.append("slides")
                n = int(re.search(r"生成 (\d+) 页内容页", user).group(1))
                assert n <= courseware.SLIDES_PER_CALL, f"一次要了 {n} 页，块太大"
                m = re.search(r"本次只出第 (\d+)~(\d+) 页", user)
                lo = int(m.group(1)) if m else 1
                return {"slides": [{"title": f"第{lo + i}页要点", "lines": ["要点一"],
                                    "script": "口播稿。"} for i in range(n)]}
            if "道题" in user:
                self.sent.append(user)
                self.calls.append("quiz")
                n = int(re.search(r"原文出 (\d+) 道题", user).group(1))
                assert n <= courseware.QUESTIONS_PER_CALL, f"一次要了 {n} 道题，块太大"
                return {"questions": [{"type": "blank", "stem": f"第{i}题", "options": [],
                                       "answer": "答案", "explanation": "解析"}
                                      for i in range(1, n + 1)]}
            return super().chat_json(system, user, images=images, max_tokens=max_tokens)

    out = fresh_dir(Path("test_out/chunked"))
    opts = make_opts(out, chapter_pages=(1,), quiz_count=15)
    llm = ChunkLLM()
    courseware.run(llm, opts, log=log_quiet, progress=lambda f, m: None)

    assert llm.calls.count("slides") == 3, f"7 页应拆成 3 次请求：{llm.calls}"
    script = next(out.rglob("口播稿.txt")).read_text(encoding="utf-8")
    marks = [i for i in range(1, 8) if f"[第{i}页]" in script]
    assert marks == list(range(1, 8)), f"拼接后页序不对：{marks}"
    assert llm.calls.count("quiz") == 2, f"15 题应拆成 2 次请求：{llm.calls}"
    assert sum("前面已出过的题" in u for u in llm.sent) == 1, "续块该带上已出的题防重复"
    print(f"  → 拆小块 OK（7 页 3 次请求、15 题 2 次请求，拼起来仍是 7 页 / 15 题）")


def test_truncation_salvage():
    """接口把回复截断时：救下已写完整的内容、原始回复留档，整节课不失败。"""
    from kouboppt import courseware
    from kouboppt.llm import LLMError

    class CutLLM(StubLLM):
        def chat_json(self, system, user, images=None, max_tokens=0, **kw):
            if "内容页" in user and "大纲" not in user:      # 大纲请求里也有"内容页"三个字

                self.sent.append(user)
                self.calls.append("slides")
                exc = LLMError("JSON 解析失败：Expecting ':' delimiter: line 1 column 2048")
                exc.raw_text = ('{"slides":[{"title":"救回来的第一页","lines":["要点"],'
                                '"script":"稿。"},{"title":"尾巴","lines":["要点')
                raise exc
            return super().chat_json(system, user, images=images, max_tokens=max_tokens)

    out = fresh_dir(Path("test_out/salvage"))
    opts = make_opts(out, chapter_pages=(1,), gen_quiz=False)
    llm = CutLLM()
    logs: list[str] = []
    llm.log = logs.append
    outs = courseware.run(llm, opts, log=logs.append, progress=lambda f, m: None)

    assert any("抢救" in s for s in logs), logs[-6:]
    assert list(out.rglob("*PPT.pptx")), "截断后仍该出课件（用救回来的内容）"
    assert "救回来的第一页" in next(out.rglob("口播稿.txt")).read_text(encoding="utf-8")
    assert list(out.rglob("AI原始回复.log")), "没留原始回复存档"
    assert any("PPT.pptx" in p.name for p in outs)
    print("  → 截断抢救 OK（每小节救回 1 页、写 AI原始回复.log、整节课照常产出）")


def test_video_overlap():
    """第一步内部流水线：第 1 节的视频在合成时，后面几节的课件已经在生成。"""
    from kouboppt import courseware, pipeline, slidegen

    out = fresh_dir(Path("test_out/overlap_video"))
    opts = make_opts(out, chapter_pages=(1, 1, 1), gen_video=True,
                     gen_notes=False, gen_quiz=False)
    llm = SlowLLM(delay=0.1)
    events: list[tuple[str, int]] = []
    lock = threading.Lock()
    orig_build = slidegen.BuiltinEngine.build
    orig_proc = pipeline.Pipeline.process_file

    def fake_build(self, lesson, path):
        orig_build(self, lesson, path)
        with lock:
            events.append(("课件", idx_of(lesson.title)))

    def fake_process_file(self, ppt_path, ranges_spec, out_dir, cancel=None, out_base=None):
        assert out_base, "教材→课程该给视频指定 out_base（课时标题+视频）"
        idx = idx_of(out_base)
        with lock:
            events.append(("视频开始", idx))
        time.sleep(0.6)
        Path(out_dir).mkdir(parents=True, exist_ok=True)
        mp4 = Path(out_dir) / f"{out_base or Path(ppt_path).stem}.mp4"
        mp4.write_bytes(b"mp4")
        with lock:
            events.append(("视频完成", idx))
        return [mp4]

    slidegen.BuiltinEngine.build = fake_build
    pipeline.Pipeline.process_file = fake_process_file
    try:
        outs = courseware.run(llm, opts, log=log_quiet, progress=lambda f, m: None)
    finally:
        slidegen.BuiltinEngine.build = orig_build
        pipeline.Pipeline.process_file = orig_proc

    assert [i for k, i in events if k == "视频完成"] == [1, 2, 3], f"视频没按书序做：{events}"
    starts = {i: n for n, (k, i) in enumerate(events) if k == "视频开始"}
    builds = {i: n for n, (k, i) in enumerate(events) if k == "课件"}
    assert starts[1] > builds[1], f"视频该在课件之后开始：{events}"
    assert starts[1] < builds[3], f"第 1 节视频没和后面几节的课件重叠：{events}"
    assert len(list(out.rglob("*.mp4"))) == 3
    assert sum("视频.mp4" in p.name for p in outs) == 3, outs
    print(f"  → 课件/视频流水线 OK（第 1 节视频合成期间，第 2、3 节课件已并行生成，"
          f"视频仍按书序出）")


def test_video_defer_to_notes():
    """全自动：课件一做完就开始笔记/题库，视频在后台继续跑；分步确认仍在闸门前收完视频。"""
    from kouboppt import courseware, docgen, pipeline

    orig_proc = pipeline.Pipeline.process_file
    orig_md = docgen.markdown_to_docx

    def run_once(gated: bool):
        out = fresh_dir(Path("test_out/defer_video_gated" if gated
                             else "test_out/defer_video_auto"))
        opts = make_opts(out, chapter_pages=(1, 1, 1), gen_video=True,
                         gen_notes=True, gen_quiz=False)
        events: list[tuple[str, int]] = []
        lock = threading.Lock()

        def rec(tag: str, idx: int):
            with lock:
                events.append((tag, idx))

        def fake_process_file(self, ppt_path, ranges_spec, out_dir, cancel=None, out_base=None):
            assert out_base, "教材→课程该给视频指定 out_base（课时标题+视频）"
            idx = idx_of(out_base)
            rec("视频开始", idx)
            time.sleep(0.8)
            Path(out_dir).mkdir(parents=True, exist_ok=True)
            mp4 = Path(out_dir) / f"{out_base or Path(ppt_path).stem}.mp4"
            mp4.write_bytes(b"mp4")
            rec("视频完成", idx)
            return [mp4]

        def fake_md(md, path, book_title="", chapter="", formula_dir=None, title=""):
            p = orig_md(md, path, book_title, chapter, formula_dir=formula_dir)
            if not title:              # 课后习题文档不算"笔记"时点
                rec("笔记", idx_of(chapter))   # 课时标题形如「第1课 p1-p1」
            return p

        pipeline.Pipeline.process_file = fake_process_file
        docgen.markdown_to_docx = fake_md
        try:
            outs = courseware.run(SlowLLM(delay=0.05), opts, log=log_quiet,
                                  progress=lambda f, m: None,
                                  gate=(lambda *a: True) if gated else None)
        finally:
            pipeline.Pipeline.process_file = orig_proc
            docgen.markdown_to_docx = orig_md
        return events, outs, out

    def spans_of(events):
        starts, spans = {}, []
        for n, (tag, idx) in enumerate(events):
            if tag == "视频开始":
                starts[idx] = n
            elif tag == "视频完成":
                spans.append((starts[idx], n))
        return spans

    # ---- 全自动：笔记与视频重叠，但视频一个都不能少
    ev, outs, out = run_once(gated=False)
    assert [i for t, i in ev if t == "视频完成"] == [1, 2, 3], f"视频没按书序出：{ev}"
    assert sorted(i for t, i in ev if t == "笔记") == [1, 2, 3], f"笔记缺节：{ev}"
    sp = spans_of(ev)
    par = [n for n, (t, _) in enumerate(ev) if t == "笔记" and any(s < n < e for s, e in sp)]
    assert par, f"全自动下笔记没有和视频并行：{ev}"
    assert len(list(out.rglob("*.mp4"))) == 3 and len(list(out.rglob("*学习笔记.docx"))) == 3
    assert sum(p.suffix == ".mp4" for p in outs) == 3, outs
    assert sum("学习笔记" in p.name for p in outs) == 3, outs

    # ---- 分步确认：闸门前必须把视频收干净（用户点"就到这里"时后台不能还在写文件）
    ev2, _, out2 = run_once(gated=True)
    assert [i for t, i in ev2 if t == "视频完成"] == [1, 2, 3], f"视频没按书序出：{ev2}"
    last_video = max(n for n, (t, _) in enumerate(ev2) if t == "视频完成")
    first_note = min(n for n, (t, _) in enumerate(ev2) if t == "笔记")
    assert last_video < first_note, f"分步模式没等视频收尾就进了第二步：{ev2}"
    assert len(list(out2.rglob("*.mp4"))) == 3
    print(f"  → 视频收尾时机 OK（全自动：笔记与后台视频重叠 {len(par)} 次；"
          f"分步确认：3 个视频全部收完才进第二步）")


def test_video_stops_on_error():
    """第二步炸了的时候，后台视频要立刻收到停止信号，别留线程继续往课节目录写文件。"""
    from kouboppt import courseware, docgen, pipeline

    out = fresh_dir(Path("test_out/defer_video_error"))
    opts = make_opts(out, chapter_pages=(1, 1, 1), gen_video=True,
                     gen_notes=True, gen_quiz=False)
    orig_proc = pipeline.Pipeline.process_file
    orig_md = docgen.markdown_to_docx
    lock = threading.Lock()
    started: list[int] = []
    saw_stop: list[tuple[int, float]] = []

    def fake_process_file(self, ppt_path, ranges_spec, out_dir, cancel=None, out_base=None):
        assert out_base, "教材→课程该给视频指定 out_base（课时标题+视频）"
        idx = idx_of(out_base)
        with lock:
            started.append(idx)
        t0 = time.monotonic()
        while time.monotonic() - t0 < 5:
            if cancel is not None and cancel.is_set():
                with lock:
                    saw_stop.append((idx, time.monotonic() - t0))
                raise pipeline.Cancelled("已取消")
            time.sleep(0.05)
        Path(out_dir).mkdir(parents=True, exist_ok=True)
        mp4 = Path(out_dir) / f"{out_base or Path(ppt_path).stem}.mp4"
        mp4.write_bytes(b"mp4")
        return [mp4]

    def boom(md, path, book_title="", chapter="", formula_dir=None, title=""):
        raise RuntimeError("笔记生成炸了（模拟）")

    pipeline.Pipeline.process_file = fake_process_file
    docgen.markdown_to_docx = boom
    try:
        courseware.run(SlowLLM(delay=0.05), opts, log=log_quiet, progress=lambda f, m: None)
        raise AssertionError("第二步出错应该抛出来")
    except RuntimeError as exc:
        assert "炸了" in str(exc), exc
    finally:
        pipeline.Pipeline.process_file = orig_proc
        docgen.markdown_to_docx = orig_md

    # 后台线程每 0.05 秒才轮询一次取消，给它一点时间注意到
    t_wait = time.monotonic()
    while not saw_stop and time.monotonic() - t_wait < 3:
        time.sleep(0.02)

    assert saw_stop, f"后台视频没收到停止信号：started={started}"
    assert saw_stop[0][1] < 3.0, f"停止信号来得太慢：{saw_stop}"
    assert len(started) < 3, f"排队中的视频没被丢弃：{started}"
    print(f"  → 出错收尾 OK（第 {saw_stop[0][0]} 节视频 {saw_stop[0][1]:.2f} 秒内停下，"
          f"只启动了 {len(started)} 节，排队中的已丢弃）")


def test_output_layout_and_legacy_move():
    """课件/视频/笔记/题库各自进独立文件夹，文件名 = 课时标题 + 类型；旧结构产物原地搬过去继续当缓存。"""
    from kouboppt import courseware, pipeline

    out = fresh_dir(Path("test_out/layout"))
    opts = make_opts(out, chapter_pages=(1,), gen_video=True)
    book_dir = out / "测试教材"
    st_dir = book_dir / "01_第1课_p1-p1"
    st_dir.mkdir(parents=True)
    (st_dir / "教材原文.md").write_text("第 1 页 扫描内容 $x$。" * 60, encoding="utf-8")
    (st_dir / "口播稿.txt").write_text("[第1页] 旧讲稿\n", encoding="utf-8")
    (st_dir / "01_第1课_p1-p1_课件.pptx").write_bytes(b"pptx")
    (st_dir / "01_第1课_p1-p1_课件.mp4").write_bytes(b"mp4")
    (st_dir / "01_学习笔记.docx").write_bytes(b"docx")
    (st_dir / "01_题库.docx").write_bytes(b"docx")

    orig = pipeline.Pipeline.process_file

    def boom(*a, **kw):
        raise AssertionError("旧产物该被当缓存复用，不该真渲染视频")

    pipeline.Pipeline.process_file = boom
    llm = StubLLM()
    try:
        outs = courseware.run(llm, opts, log=log_quiet, progress=lambda f, m: None)
    finally:
        pipeline.Pipeline.process_file = orig

    assert not llm.calls, f"全是缓存，一次 AI 都不该调：{llm.calls}"
    ppt = book_dir / courseware.PPT_DIR / "第1课_p1-p1PPT.pptx"
    vid = book_dir / courseware.VIDEO_DIR / "第1课_p1-p1视频.mp4"
    note = book_dir / courseware.NOTE_DIR / "第1课_p1-p1学习笔记.docx"
    quiz = book_dir / courseware.QUIZ_DIR / "第1课_p1-p1题库.docx"
    assert ppt.exists() and vid.exists(), [str(p) for p in book_dir.rglob("*")]
    assert note.exists() and quiz.exists(), [str(p) for p in book_dir.rglob("*")]
    assert not (st_dir / "01_第1课_p1-p1_课件.pptx").exists(), "旧课件该搬走"
    assert not (st_dir / "01_第1课_p1-p1_课件.mp4").exists(), "旧视频该搬走"
    assert not (st_dir / "01_学习笔记.docx").exists(), "旧笔记该搬走"
    assert not (st_dir / "01_题库.docx").exists(), "旧题库该搬走"
    assert ppt in outs and vid in outs and note in outs and quiz in outs, outs
    assert (st_dir / "教材原文.md").exists(), "教材原文.md 不动"
    assert (st_dir / "口播稿.txt").exists(), "口播稿.txt 不动"
    print(f"  → 新布局 OK（{courseware.PPT_DIR}/、{courseware.VIDEO_DIR}/、"
          f"{courseware.NOTE_DIR}/、{courseware.QUIZ_DIR}/ 各就各位，"
          "旧结构原地搬迁后照样算缓存）")


def test_exercises_extracted():
    """课后习题/思考题：单独出 课后习题/课时标题课后习题.docx，且不进课件/笔记/题库请求。"""
    from kouboppt import courseware
    from docx import Document

    body_line = "直角三角形的两条直角边的平方和等于斜边的平方，这就是勾股定理。\n"
    ex_md = ("# 第1章 测试章节\n" + body_line * 8 +
             "## 课后习题\n1. 已知两直角边 3、4，求斜边的长度。\n"
             "2. 试证明勾股定理的逆定理也成立，并举例说明。\n"
             "## 思考题\n3. 想一想：商高定理这个名字是怎么来的？查资料说明。\n")

    class ExLLM(StubLLM):
        """识别稿里带课后习题（正文够长，能过"识别内容过少"的判据）。"""

        def chat(self, system, user, images=None, max_tokens=0, **kw):
            if "转写为 Markdown" in user:
                self.sent.append(user)
                self.calls.append("ocr")
                return ex_md
            if "学习笔记" in user:
                self.sent.append(user)
                self.calls.append("notes")
                return "# 知识框架\n- 勾股定理 $a^2+b^2=c^2$"
            raise AssertionError("桩 LLM 没预料到的 chat 调用：" + user[:80])

    out = fresh_dir(Path("test_out/exercises"))
    opts = make_opts(out, chapter_pages=(1,))
    make_scanned_pdf(opts.pdf_path, 1)          # 带 $ → 走识别路径，识别结果由桩给出
    llm = ExLLM()
    outs = courseware.run(llm, opts, log=log_quiet, progress=lambda f, m: None)

    book_dir = out / "测试教材"
    ex_doc = book_dir / courseware.EX_DIR / "第1课_p1-p1课后习题.docx"
    assert ex_doc.exists(), [str(p) for p in book_dir.rglob("*")]
    assert ex_doc in outs, outs
    text = "\n".join(p.text for p in Document(str(ex_doc)).paragraphs)
    assert "课后习题" in text and "逆定理" in text and "商高定理这个名字" in text, text
    assert body_line.strip() not in text, "正文不该混进课后习题文档"

    # 教材原文.md 保留全文（含习题），但习题不进课件/笔记/题库的请求
    md_txt = next(book_dir.rglob("教材原文.md")).read_text(encoding="utf-8")
    assert "## 课后习题" in md_txt, "教材原文.md 该保留全文"
    for kw, name in (("大纲", "课件大纲"), ("内容页", "小节内容"),
                     ("学习笔记", "学习笔记"), ("道题", "题库")):
        hits = [u for u in llm.sent if kw in u]
        assert hits, f"没发出「{name}」请求"
        assert all("逆定理" not in u and "商高定理这个名字" not in u for u in hits), \
            f"「{name}」请求里混进了课后习题内容"
        assert all("勾股定理" in u for u in hits), f"「{name}」请求里没带教材原文"
    assert llm.calls.count("notes") == 2, llm.calls       # 笔记仍分上下篇两次请求

    # 重跑：课后习题、课件、笔记全算缓存，一次 AI 都不调
    mtime = ex_doc.stat().st_mtime
    llm2 = ExLLM()
    outs2 = courseware.run(llm2, opts, log=log_quiet, progress=lambda f, m: None)
    assert not llm2.calls, f"全是缓存，一次 AI 都不该调：{llm2.calls}"
    assert ex_doc.stat().st_mtime == mtime and ex_doc in outs2, outs2
    print("  → 课后习题单独成文 OK（Word 在 课后习题/，不进课件与笔记，重跑算缓存）")


def main():
    print("== 1/20 全自动（按每节页数）==")
    test_auto()
    print("== 2/20 自动选风格（AI 按书名 + 缓存）==")
    test_auto_style()
    print("== 3/20 公式兜底两层（改写重渲 / 原页裁图）==")
    test_formula_fallback()
    print("== 4/20 公式兜底端到端（坏公式进 PPT + 缓存）==")
    test_formula_fallback_e2e()
    print("== 5/20 分步 + 续跑 ==")
    test_step_and_resume()
    print("== 6/20 只出笔记 ==")
    test_notes_only()
    print("== 7/20 按教材章节 + 缓存 ==")
    test_chapter_mode()
    print("== 8/20 章节内不跨章 ==")
    test_no_cross_chapter()
    print("== 9/20 识别页级并发 ==")
    test_ocr_concurrency()
    print("== 10/20 课件小节并发 ==")
    test_lesson_concurrency()
    print("== 11/20 笔记/题库并发 ==")
    test_notes_quiz_concurrency()
    print("== 12/20 视频页级并发 ==")
    test_video_page_concurrency()
    print("== 13/20 提示词都带教材原文 ==")
    test_prompts_carry_source()
    print("== 14/20 拆小块防截断 ==")
    test_chunked_requests()
    print("== 15/20 截断抢救 ==")
    test_truncation_salvage()
    print("== 16/20 课件/视频流水线并行 ==")
    test_video_overlap()
    print("== 17/20 全自动下视频压到第二步之后收尾 ==")
    test_video_defer_to_notes()
    print("== 18/20 第二步出错时叫停后台视频 ==")
    test_video_stops_on_error()
    print("== 19/20 产物分类文件夹 + 旧结构搬迁 ==")
    test_output_layout_and_legacy_move()
    print("== 20/20 课后习题单独成文 ==")
    test_exercises_extracted()
    print("COURSEWARE OK")



if __name__ == "__main__":
    main()
